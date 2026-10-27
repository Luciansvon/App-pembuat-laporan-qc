"""Script to generate a complete, realistic example of a Combined QC Report using the Inspectra application."""
from __future__ import annotations

import io
import math
import os
import shutil
import sys
from pathlib import Path

# Add project root and backend to sys.path
root_dir = Path(__file__).resolve().parents[1]
backend_dir = root_dir / "backend"
if str(backend_dir) not in sys.path:
    sys.path.insert(0, str(backend_dir))
if str(root_dir) not in sys.path:
    sys.path.insert(0, str(root_dir))

from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont

from app.core.config import Settings
from app.main import create_app


def get_font(size: int = 20, bold: bool = False):
    try:
        font_name = "segoeuib.ttf" if bold else "segoeui.ttf"
        return ImageFont.truetype(font_name, size)
    except Exception:
        try:
            font_name = "arialbd.ttf" if bold else "arial.ttf"
            return ImageFont.truetype(font_name, size)
        except Exception:
            return ImageFont.load_default()


def create_base_canvas(width: int = 1024, height: int = 768, bg_color: str = "#f8fafc") -> tuple[Image.Image, ImageDraw.ImageDraw]:
    img = Image.new("RGB", (width, height), bg_color)
    draw = ImageDraw.Draw(img)
    # Outer border
    draw.rectangle([10, 10, width - 11, height - 11], outline="#cbd5e1", width=3)
    return img, draw


def add_banner(draw: ImageDraw.ImageDraw, section: str, title: str, subtitle: str = ""):
    # Header bar
    draw.rectangle([12, 12, 1012, 85], fill="#1e293b")
    draw.rectangle([25, 24, 180, 52], fill="#3b82f6")
    font_badge = get_font(16, bold=True)
    draw.text((35, 28), "INSPECTRA QC", fill="#ffffff", font=font_badge)

    font_sec = get_font(18, bold=True)
    draw.text((200, 27), f"SECTION: {section}", fill="#93c5fd", font=font_sec)

    font_sub = get_font(15, bold=False)
    draw.text((200, 56), title + (f" | {subtitle}" if subtitle else ""), fill="#e2e8f0", font=font_sub)

    # Footer bar
    draw.rectangle([12, 715, 1012, 756], fill="#334155")
    font_foot = get_font(14, bold=False)
    draw.text((25, 725), "PO 248 | ERICE STACKING CHAIR | CUSTOMER: RODA | DATE: 2026-10-21", fill="#cbd5e1", font=font_foot)
    draw.text((820, 725), "QC: BIMA QC [PASS]", fill="#4ade80", font=get_font(14, bold=True))


def render_cardboard_box(draw: ImageDraw.ImageDraw, x: int, y: int, w: int, h: int, text_label: str = ""):
    # Box front
    draw.rectangle([x, y, x + w, y + h], fill="#d97706", outline="#92400e", width=4)
    # Box texture fluting lines
    for ly in range(y + 20, y + h - 10, 25):
        draw.line([(x + 10, ly), (x + w - 10, ly)], fill="#b45309", width=1)
    # Center packaging tape
    draw.rectangle([x + w // 2 - 25, y, x + w // 2 + 25, y + h], fill="#fed7aa", outline="#ca8a04", width=2)
    # Shipping mark label
    lw, lh = 180, 110
    lx, ly = x + 30, y + 40
    draw.rectangle([lx, ly, lx + lw, ly + lh], fill="#ffffff", outline="#64748b", width=2)
    font_lbl = get_font(13, bold=True)
    draw.text((lx + 10, ly + 8), "CUSTOMER: RODA", fill="#0f172a", font=font_lbl)
    draw.text((lx + 10, ly + 28), "ITEM: ERICE CHAIR", fill="#0f172a", font=font_lbl)
    draw.text((lx + 10, ly + 48), "PO: 248", fill="#0f172a", font=font_lbl)
    # Barcode
    for bx in range(lx + 10, lx + lw - 20, 5):
        draw.line([(bx, ly + 72), (bx, ly + 98)], fill="#0f172a", width=2)

    # Box label text
    if text_label:
        font_txt = get_font(18, bold=True)
        draw.rectangle([x + w // 2 - 140, y + h - 60, x + w // 2 + 140, y + h - 15], fill="#0f172a")
        draw.text((x + w // 2 - 120, y + h - 52), text_label, fill="#ffffff", font=font_txt)


def render_chair(draw: ImageDraw.ImageDraw, cx: int, cy: int, scale: float = 1.0):
    # Draw ergonomic stacking chair wireframe/render
    frame_color = "#1e293b"
    wood_color = "#ca8a04"

    # Legs (tubular)
    fl_left = (int(cx - 120 * scale), int(cy + 180 * scale))
    fl_right = (int(cx + 120 * scale), int(cy + 180 * scale))
    bl_left = (int(cx - 150 * scale), int(cy + 220 * scale))
    bl_right = (int(cx + 150 * scale), int(cy + 220 * scale))

    # Seat corners
    s_fl = (int(cx - 100 * scale), int(cy + 20 * scale))
    s_fr = (int(cx + 100 * scale), int(cy + 20 * scale))
    s_bl = (int(cx - 90 * scale), int(cy - 40 * scale))
    s_br = (int(cx + 90 * scale), int(cy - 40 * scale))

    # Backrest top
    b_tl = (int(cx - 90 * scale), int(cy - 180 * scale))
    b_tr = (int(cx + 90 * scale), int(cy - 180 * scale))

    # Draw legs
    draw.line([s_fl, fl_left], fill=frame_color, width=int(10 * scale))
    draw.line([s_fr, fl_right], fill=frame_color, width=int(10 * scale))
    draw.line([s_bl, bl_left], fill=frame_color, width=int(10 * scale))
    draw.line([s_br, bl_right], fill=frame_color, width=int(10 * scale))

    # Nylon glides
    for pt in [fl_left, fl_right, bl_left, bl_right]:
        draw.ellipse([pt[0] - 10, pt[1] - 5, pt[0] + 10, pt[1] + 12], fill="#475569")

    # Seat slats
    draw.polygon([s_bl, s_br, s_fr, s_fl], fill=wood_color, outline="#78350f")
    for i in range(1, 5):
        y_off = s_bl[1] + (s_fl[1] - s_bl[1]) * i // 5
        draw.line([(s_bl[0] + 10, y_off), (s_br[0] - 10, y_off)], fill="#78350f", width=2)

    # Backrest frame & slats
    draw.line([s_bl, b_tl], fill=frame_color, width=int(10 * scale))
    draw.line([s_br, b_tr], fill=frame_color, width=int(10 * scale))
    draw.line([b_tl, b_tr], fill=frame_color, width=int(12 * scale))

    # Backrest teak slats
    for i in range(1, 4):
        sy = b_tl[1] + 35 * i
        draw.rounded_rectangle([b_tl[0] + 15, sy, b_tr[0] - 15, sy + 22], radius=4, fill=wood_color, outline="#78350f")

    # Armrests
    draw.line([(b_tl[0], b_tl[1] + 90), (s_fl[0], s_fl[1] - 40)], fill=frame_color, width=int(8 * scale))
    draw.line([(b_tr[0], b_tr[1] + 90), (s_fr[0], s_fr[1] - 40)], fill=frame_color, width=int(8 * scale))


def render_tape_measure(draw: ImageDraw.ImageDraw, x1: int, y1: int, x2: int, y2: int, value_mm: int, label: str):
    # Yellow measuring tape
    dx = x2 - x1
    dy = y2 - y1
    dist = math.hypot(dx, dy)
    angle = math.atan2(dy, dx)

    # Draw tape line
    tape_color = "#facc15"
    draw.line([(x1, y1), (x2, y2)], fill=tape_color, width=28)
    draw.line([(x1, y1), (x2, y2)], fill="#ca8a04", width=30)
    draw.line([(x1, y1), (x2, y2)], fill=tape_color, width=24)

    # Ticks along the line
    num_ticks = max(5, int(dist // 25))
    for i in range(num_ticks + 1):
        t = i / num_ticks
        tx = x1 + dx * t
        ty = y1 + dy * t
        # Perpendicular
        px = -dy / dist * (10 if i % 2 == 0 else 6)
        py = dx / dist * (10 if i % 2 == 0 else 6)
        draw.line([(tx - px, ty - py), (tx + px, ty + py)], fill="#0f172a", width=2)

    # Callout badge at midpoint
    mx = (x1 + x2) // 2
    my = (y1 + y2) // 2
    bw, bh = 220, 52
    draw.rounded_rectangle([mx - bw // 2, my - bh // 2, mx + bw // 2, my + bh // 2], radius=8, fill="#0f172a", outline="#facc15", width=3)
    draw.text((mx - bw // 2 + 16, my - 16), f"{label}: {value_mm} mm", fill="#ffffff", font=get_font(18, bold=True))


def render_macro_tape(draw: ImageDraw.ImageDraw, value_mm: int, label: str):
    # Macro zoomed-in view of tape measure
    draw.rectangle([50, 140, 974, 650], fill="#f1f5f9", outline="#64748b", width=3)
    font_hdr = get_font(24, bold=True)
    draw.text((80, 165), f"DETAIL PENGUKURAN PRESISI — {label}", fill="#0f172a", font=font_hdr)

    # Huge yellow tape band
    draw.rectangle([80, 280, 944, 460], fill="#fde047", outline="#ca8a04", width=4)

    # Big millimeter tick marks
    center_x = 512
    for offset in range(-8, 9):
        x = center_x + offset * 45
        tick_h = 100 if offset == 0 else (60 if offset % 5 == 0 else 35)
        color = "#dc2626" if offset == 0 else "#0f172a"
        width = 5 if offset == 0 else (3 if offset % 5 == 0 else 2)
        draw.line([(x, 460 - tick_h), (x, 460)], fill=color, width=width)
        if offset % 2 == 0:
            val = value_mm + offset
            font_t = get_font(18, bold=(offset == 0))
            draw.text((x - 18, 305), str(val), fill=color, font=font_t)

    # Red needle / pointer
    draw.polygon([(center_x, 490), (center_x - 18, 540), (center_x + 18, 540)], fill="#dc2626")
    draw.rectangle([center_x - 140, 555, center_x + 140, 615], fill="#dc2626")
    draw.text((center_x - 120, 568), f"ACTUAL: {value_mm} mm", fill="#ffffff", font=get_font(22, bold=True))


def generate_all_sample_photos() -> dict[str, list[tuple[str, bytes]]]:
    photos: dict[str, list[tuple[str, bytes]]] = {}

    # 1. PACKING_ATTRIBUTE (6 photos)
    attr_list = [
        ("Master Box Shipping Mark & Barcode PO 248", "Pemeriksaan label luar kardus, barcode jelas terbaca"),
        ("Logo RODA pada Sisi Luar Karton", "Cetak logo RODA rapi dan presisi di tengah"),
        ("Label QC Passed & Hologram Pabrik", "Verifikasi stiker kelulusan standar ekspor"),
        ("Silica Gel & Kantong Pengering 50g", "Pencegah kelembapan di dalam kemasan"),
        ("Busa EPE Pelindung Sudut Kaki", "Ketebalan busa 15 mm pelindung benturan"),
        ("Instruksi Penanganan Fragile & Keep Dry", "Simbol penanganan barang pecah belah lengkap"),
    ]
    photos["PACKING_ATTRIBUTE"] = []
    for i, (title, desc) in enumerate(attr_list, 1):
        img, draw = create_base_canvas()
        add_banner(draw, "PACKING ATTRIBUTE", title, desc)
        render_cardboard_box(draw, 212, 140, 600, 520, title)
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=92)
        photos["PACKING_ATTRIBUTE"].append((f"attr_{i}.jpg", buf.getvalue()))

    # 2. BOX_DIMENSION (3 axes x 2 photos = 6 photos)
    box_axes = [("Length", 770, "Panjang Kardus Luar"), ("Width", 680, "Lebar Kardus Luar"), ("Height", 945, "Tinggi Kardus Luar")]
    photos["BOX_DIMENSION"] = []
    for axis, val, desc in box_axes:
        # Slot 1: Overview
        img1, draw1 = create_base_canvas()
        add_banner(draw1, "BOX DIMENSION", f"PENGUKURAN {axis.upper()} (OVERVIEW)", f"Nilai Terukur: {val} mm")
        render_cardboard_box(draw1, 262, 160, 500, 480)
        if axis == "Length":
            render_tape_measure(draw1, 262, 600, 762, 600, val, "Length (L)")
        elif axis == "Width":
            render_tape_measure(draw1, 262, 380, 762, 380, val, "Width (W)")
        else:
            render_tape_measure(draw1, 720, 160, 720, 640, val, "Height (H)")
        b1 = io.BytesIO()
        img1.save(b1, format="JPEG", quality=92)
        photos["BOX_DIMENSION"].append((f"box_{axis}_overview.jpg", b1.getvalue()))

        # Slot 2: Macro Close-up
        img2, draw2 = create_base_canvas()
        add_banner(draw2, "BOX DIMENSION", f"PENGUKURAN {axis.upper()} (CLOSE-UP METERAN)", f"Tepat di Garis {val} mm")
        render_macro_tape(draw2, val, f"BOX {axis.upper()}")
        b2 = io.BytesIO()
        img2.save(b2, format="JPEG", quality=92)
        photos["BOX_DIMENSION"].append((f"box_{axis}_closeup.jpg", b2.getvalue()))

    # 3. PACKING_METHOD (6 photos)
    method_list = [
        ("Penyusunan Tumpukan Kursi (Stacking 4 Pcs)", "Penempatan kursi bertumpuk rapi di dalam dus"),
        ("Pemisah Karton Beralur Antar Dudukan", "Karton pelindung mencegah lecet antar dudukan"),
        ("Plastik Wrapping Selubung Produk", "Plastik transparan kedap debu dan air"),
        ("Bantalan Bawah Kardus (Bottom Cushion)", "Karton tebal 5 ply alas dasar box"),
        ("Penutupan Tutup Box (Top Fold)", "Lipatan atas box rata dan rapat"),
        ("Penyegelan Lakban Bening H-Taping", "Segel lakban kuat melintang dan membujur"),
    ]
    photos["PACKING_METHOD"] = []
    for i, (title, desc) in enumerate(method_list, 1):
        img, draw = create_base_canvas()
        add_banner(draw, "PACKING METHOD", title, desc)
        render_cardboard_box(draw, 162, 160, 700, 500)
        render_chair(draw, 512, 400, scale=0.7)
        font_m = get_font(20, bold=True)
        draw.rectangle([250, 600, 774, 650], fill="#047857")
        draw.text((270, 612), f"METODE: {title}", fill="#ffffff", font=font_m)
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=92)
        photos["PACKING_METHOD"].append((f"method_{i}.jpg", buf.getvalue()))

    # 4. PRODUCT_VIEW (4 photos)
    views = [
        ("Tampak Depan (Front View)", "Bentuk simetris, dudukan rata, rangka kokoh"),
        ("Tampak Samping (Side View)", "Kemiringan sandaran 105 derajat sesuai ergonomi"),
        ("Tampak Belakang (Back View)", "Lekukan pipa sandaran halus tanpa penyok"),
        ("Tampak Isometrik 3D (Angle View)", "Kombinasi rangka powder coat hitam & kayu jati"),
    ]
    photos["PRODUCT_VIEW"] = []
    for i, (title, desc) in enumerate(views, 1):
        img, draw = create_base_canvas(bg_color="#f1f5f9")
        add_banner(draw, "PRODUCT VIEW", title, desc)
        render_chair(draw, 512, 420, scale=1.1)
        font_v = get_font(22, bold=True)
        draw.rectangle([260, 640, 764, 695], fill="#1e293b")
        draw.text((285, 652), title, fill="#ffffff", font=font_v)
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=92)
        photos["PRODUCT_VIEW"].append((f"view_{i}.jpg", buf.getvalue()))

    # 5. PRODUCT_DETAIL (4 photos)
    details = [
        ("Detail Sambungan Las TIG Rangka Aluminium", "Hasil las rapi, matang sempurna, dan di-finishing rata"),
        ("Kaki Kursi (Nylon Glide Protector)", "Plug nilon hitam terpasang presisi tanpa gores"),
        ("Bumper Karet Dudukan Anti Benturan Tumpukan", "Karet elastis peredam benturan antar kursi"),
        ("Logo Timbul Gravir Laser RODA", "Posisi logo presisi di rangka belakang sesuai standar"),
    ]
    photos["PRODUCT_DETAIL"] = []
    for i, (title, desc) in enumerate(details, 1):
        img, draw = create_base_canvas(bg_color="#e2e8f0")
        add_banner(draw, "PRODUCT DETAIL", title, desc)
        # Macro component representation
        draw.rectangle([150, 160, 874, 620], fill="#334155", outline="#64748b", width=4)
        draw.ellipse([432, 310, 592, 470], fill="#0284c7", outline="#38bdf8", width=6)
        draw.text((200, 540), title, fill="#f8fafc", font=get_font(20, bold=True))
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=92)
        photos["PRODUCT_DETAIL"].append((f"detail_{i}.jpg", buf.getvalue()))

    # 6. DRAIN_HOLE (2 photos)
    drain_holes = [
        ("Lubang Pembuangan Air (Drain Hole) Kaki Kanan Bawah", "Diameter 5 mm, bersih tanpa gerigi logam (deburred)"),
        ("Lubang Pembuangan Air (Drain Hole) Kaki Kiri Bawah", "Aliran air lancar untuk kursi luar ruangan (outdoor)"),
    ]
    photos["DRAIN_HOLE"] = []
    for i, (title, desc) in enumerate(drain_holes, 1):
        img, draw = create_base_canvas()
        add_banner(draw, "DRAIN HOLE", title, desc)
        draw.rectangle([200, 180, 824, 600], fill="#1e293b", outline="#475569", width=4)
        # Metallic tube
        draw.rectangle([350, 250, 674, 520], fill="#475569", outline="#94a3b8", width=3)
        # Drain hole
        draw.ellipse([482, 355, 542, 415], fill="#0f172a", outline="#38bdf8", width=4)
        # Red callout arrow
        draw.line([(512, 280), (512, 345)], fill="#ef4444", width=6)
        draw.polygon([(512, 355), (500, 335), (524, 335)], fill="#ef4444")
        draw.rectangle([300, 210, 724, 260], fill="#ef4444")
        draw.text((320, 222), "DRAIN HOLE Ø 5 MM - PASS", fill="#ffffff", font=get_font(18, bold=True))
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=92)
        photos["DRAIN_HOLE"].append((f"drain_{i}.jpg", buf.getvalue()))

    # 7. DRAWING (1 photo)
    img_draw, d_draw = create_base_canvas(bg_color="#1e3a8a")
    add_banner(d_draw, "DRAWING PRODUCT", "Gambar Kerja Resmi Erice Stacking Chair", "Spesifikasi W: 587 mm x D: 572 mm x H: 764 mm")
    # Blueprint grid
    for gx in range(50, 974, 40):
        d_draw.line([(gx, 120), (gx, 690)], fill="#2563eb", width=1)
    for gy in range(120, 690, 40):
        d_draw.line([(50, gy), (974, gy)], fill="#2563eb", width=1)
    # Drawing border
    d_draw.rectangle([50, 120, 974, 690], outline="#ffffff", width=2)
    render_chair(d_draw, 512, 410, scale=1.0)
    # Dimension markings in blueprint style
    font_bp = get_font(18, bold=True)
    d_draw.text((430, 210), "H: 764 mm", fill="#ffffff", font=font_bp)
    d_draw.text((250, 590), "W: 587 mm", fill="#ffffff", font=font_bp)
    d_draw.text((680, 590), "D: 572 mm", fill="#ffffff", font=font_bp)
    buf_d = io.BytesIO()
    img_draw.save(buf_d, format="JPEG", quality=92)
    photos["DRAWING"] = [("drawing_erice.jpg", buf_d.getvalue())]

    # 8. DIMENSION (3 axes x 2 photos = 6 photos)
    prod_dims = [
        ("Width:", 587, "Lebar Produk Total"),
        ("Depth:", 572, "Kedalaman Produk Total"),
        ("Height:", 764, "Tinggi Total Kursi"),
    ]
    photos["DIMENSION"] = []
    for axis, val, desc in prod_dims:
        # Overview
        img1, draw1 = create_base_canvas()
        add_banner(draw1, "PRODUCT DIMENSION", f"PENGUKURAN {axis} (OVERVIEW)", f"Nilai Terukur: {val} mm")
        render_chair(draw1, 512, 430, scale=0.95)
        if "Width" in axis:
            render_tape_measure(draw1, 380, 580, 644, 580, val, "Lebar Dudukan")
        elif "Depth" in axis:
            render_tape_measure(draw1, 410, 410, 614, 410, val, "Kedalaman")
        else:
            render_tape_measure(draw1, 640, 260, 640, 600, val, "Tinggi Total")
        b1 = io.BytesIO()
        img1.save(b1, format="JPEG", quality=92)
        photos["DIMENSION"].append((f"dim_{axis.replace(':', '')}_ov.jpg", b1.getvalue()))

        # Close-up
        img2, draw2 = create_base_canvas()
        add_banner(draw2, "PRODUCT DIMENSION", f"PENGUKURAN {axis} (CLOSE-UP METERAN)", f"Garis Presisi {val} mm")
        render_macro_tape(draw2, val, f"PRODUK {axis.upper()}")
        b2 = io.BytesIO()
        img2.save(b2, format="JPEG", quality=92)
        photos["DIMENSION"].append((f"dim_{axis.replace(':', '')}_cu.jpg", b2.getvalue()))

    # 9. LEVEL_CHECK (2 photos)
    level_list = [
        ("Uji Kerataan 4 Kaki di atas Meja Marmer Presisi", "Pengecekan rocking/goyang — celah 0 mm (sempurna)"),
        ("Pengecekan Waterpass Digital Permukaan Dudukan", "Derajat kemiringan terukur 0.0° toleransi presisi"),
    ]
    photos["LEVEL_CHECK"] = []
    for i, (title, desc) in enumerate(level_list, 1):
        img, draw = create_base_canvas()
        add_banner(draw, "LEVEL CHECK", title, desc)
        # Granite table
        draw.rectangle([100, 540, 924, 620], fill="#1e293b", outline="#475569", width=3)
        draw.text((120, 570), "GRANITE INSPECTION SURFACE TABLE — FLAT TOLERANCE ±0.05 MM", fill="#94a3b8", font=get_font(15, bold=True))
        render_chair(draw, 512, 380, scale=0.85)
        # Digital level gadget
        draw.rectangle([412, 290, 612, 350], fill="#0f172a", outline="#22c55e", width=3)
        draw.text((450, 305), "0.0° PASS", fill="#22c55e", font=get_font(22, bold=True))
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=92)
        photos["LEVEL_CHECK"].append((f"level_{i}.jpg", buf.getvalue()))

    # 10. MAL_CHECK (2 photos)
    mal_list = [
        ("Pemasangan Kursi pada Jig Mal Perakitan Pabrik", "Semua titik tumpu masuk tepat ke dalam lubang pin jig"),
        ("Inspeksi Pin Pengunci Mal Perakitan", "Tidak ada deformasi atau distorsi pada rangka"),
    ]
    photos["MAL_CHECK"] = []
    for i, (title, desc) in enumerate(mal_list, 1):
        img, draw = create_base_canvas()
        add_banner(draw, "MAL CHECK", title, desc)
        # Jig base
        draw.rectangle([150, 200, 874, 620], fill="#0284c7", outline="#0369a1", width=4)
        draw.rectangle([200, 250, 824, 570], fill="#f8fafc", outline="#64748b", width=2)
        render_chair(draw, 512, 410, scale=0.75)
        draw.rectangle([320, 140, 704, 190], fill="#0f172a")
        draw.text((345, 152), "ASSEMBLY JIG FIXTURE — 100% FIT", fill="#38bdf8", font=get_font(18, bold=True))
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=92)
        photos["MAL_CHECK"].append((f"mal_{i}.jpg", buf.getvalue()))

    # 11. RADIUS (2 photos: R44 and R10)
    radius_list = [
        ("Uji Mal Kelengkungan Radius R44 pada Pipa Sandaran", "Radius R44 pas rapat tanpa ada celah cahaya"),
        ("Uji Mal Kelengkungan Radius R10 pada Tekukan Kaki", "Radius R10 sesuai gambar kerja teknis"),
    ]
    photos["RADIUS"] = []
    for i, (title, desc) in enumerate(radius_list, 1):
        img, draw = create_base_canvas()
        add_banner(draw, "RADIUS GAUGE", title, desc)
        r_label = "R44" if i == 1 else "R10"
        # Radius gauge tool
        draw.arc([300, 200, 724, 624], 180, 270, fill="#f59e0b", width=25)
        draw.rectangle([250, 520, 774, 580], fill="#1e293b")
        draw.text((275, 536), f"PRECISION RADIUS GAUGE {r_label} — RESULT: PASS", fill="#fde047", font=get_font(20, bold=True))
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=92)
        photos["RADIUS"].append((f"radius_{i}.jpg", buf.getvalue()))

    # 12. ISSUE (2 issues: Rough and Scratch)
    # Defect 1: Rough
    img_r, draw_r = create_base_canvas(bg_color="#334155")
    add_banner(draw_r, "ISSUE TEMUAN", "Cacat Kasar Cat (Rough Surface)", "Ditemukan pada 2 unit kursi di area rangka samping")
    # Simulated rough powder coat surface
    draw_r.rectangle([100, 160, 924, 620], fill="#1e293b", outline="#64748b", width=3)
    # Pebble dots
    for px in range(480, 545, 5):
        for py in range(360, 425, 5):
            draw_r.ellipse([px, py, px + 3, py + 3], fill="#94a3b8")
    # Red circle around defect
    draw_r.ellipse([450, 330, 574, 454], outline="#ef4444", width=5)
    # Red arrow pointing to defect
    draw_r.line([(350, 260), (450, 330)], fill="#ef4444", width=6)
    draw_r.polygon([(450, 330), (430, 310), (440, 290)], fill="#ef4444")
    # Defect label badge
    draw_r.rectangle([180, 220, 360, 270], fill="#ef4444")
    draw_r.text((195, 234), "DEFECT: ROUGH", fill="#ffffff", font=get_font(18, bold=True))
    buf_r = io.BytesIO()
    img_r.save(buf_r, format="JPEG", quality=92)

    # Defect 2: Scratch
    img_s, draw_s = create_base_canvas(bg_color="#334155")
    add_banner(draw_s, "ISSUE TEMUAN", "Goresan Halus (Scratch Mark)", "Ditemukan pada 1 unit kursi akibat gesekan troly transfer")
    draw_s.rectangle([100, 160, 924, 620], fill="#1e293b", outline="#64748b", width=3)
    # Scratch line
    draw_s.line([(480, 360), (550, 410)], fill="#e2e8f0", width=3)
    # Red box around defect
    draw_s.rectangle([460, 340, 570, 430], outline="#ef4444", width=5)
    # Arrow
    draw_s.line([(360, 280), (460, 340)], fill="#ef4444", width=6)
    draw_s.rectangle([200, 240, 380, 290], fill="#ef4444")
    draw_s.text((215, 254), "DEFECT: SCRATCH", fill="#ffffff", font=get_font(18, bold=True))
    buf_s = io.BytesIO()
    img_s.save(buf_s, format="JPEG", quality=92)

    photos["ISSUE"] = [
        ("issue_rough.jpg", buf_r.getvalue()),
        ("issue_scratch.jpg", buf_s.getvalue()),
    ]

    return photos


def main():
    print("=== MEMBUAT CONTOH LAPORAN QC GABUNGAN INSPECTRA ===")
    
    # Target databases: Desktop app SQLite
    local_app_data = Path(os.environ.get("LOCALAPPDATA", str(Path.home() / "AppData" / "Local")))
    desktop_data_dir = local_app_data / "QC Report Assistant" / "data"
    desktop_data_dir.mkdir(parents=True, exist_ok=True)
    
    settings = Settings(data_dir=desktop_data_dir)
    print(f"Menggunakan penyimpanan data: {settings.data_dir}")
    
    app = create_app(settings)
    with TestClient(app) as client:
        # 1. Pastikan Customer RODA terdaftar
        cust_res = client.get("/api/customers")
        customers = cust_res.json()
    roda_cust = next((c for c in customers if c.get("code") == "RODA"), None)
    if not roda_cust:
        create_c = client.post("/api/customers", json={"code": "RODA", "name": "RODA FURNITURE"})
        assert create_c.status_code == 201, create_c.text
        roda_cust = create_c.json()
    customer_id = roda_cust["id"]
    print(f"Customer terdaftar: {roda_cust['name']} ({roda_cust['code']})")

    # 2. Pastikan Product ERICE STACKING CHAIR terdaftar
    prod_res = client.get(f"/api/products?customer_id={customer_id}")
    products = prod_res.json()
    erice_prod = next((p for p in products if "ERICE" in p.get("name", "").upper()), None)
    if not erice_prod:
        create_p = client.post("/api/products", json={
            "customer_id": customer_id,
            "name": "ERICE STACKING CHAIR",
            "standard_dimensions": {"W": 587, "D": 572, "H": 764},
        })
        assert create_p.status_code == 201, create_p.text
        erice_prod = create_p.json()
    product_id = erice_prod["id"]
    print(f"Produk terdaftar: {erice_prod['name']}")

    # 3. Buat Inspeksi Baru dengan report_type = 'COMBINED'
    insp_payload = {
        "product_id": product_id,
        "po": "248",
        "inspection_type": "FINAL",
        "report_type": "COMBINED",
        "date": "2026-10-21",
        "location": "QSF",
        "qc_name": "BIMA QC",
        "quantity": 200,
        "product_dimensions": {"W": 587, "D": 572, "H": 764},
        "box_dimensions": {"L": 770, "W": 680, "H": 945},
    }
    insp_res = client.post("/api/inspections", json=insp_payload)
    assert insp_res.status_code == 201, insp_res.text
    inspection = insp_res.json()
    inspection_id = inspection["id"]
    print(f"Inspeksi baru berhasil dibuat: {inspection['inspection_number']} (ID: {inspection_id})")

    # Update conclusion & communication status via PATCH
    patch_res = client.patch(f"/api/inspections/{inspection_id}", json={
        "revision": inspection["revision"],
        "conclusion": "Pemeriksaan menyeluruh terhadap 200 unit ERICE STACKING CHAIR telah selesai dilakukan. Seluruh item memenuhi spesifikasi gambar kerja, dimensi box, dan standar kemasan PO 248. Temuan cacat minor (kasar cat & goresan halus) telah diperbaiki di tempat dan disetujui.",
        "communication_status": "All found issue have been communicated and repaired by factory.",
    })
    assert patch_res.status_code == 200, patch_res.text

    # 4. Generate semua foto realistis
    print("Menghasilkan foto simulasi QC realistis untuk semua seksi...")
    sample_photos = generate_all_sample_photos()

    # 5. Upload Packing Attribute
    print("Mengunggah foto Packing Attribute...")
    for fname, fbytes in sample_photos["PACKING_ATTRIBUTE"]:
        res = client.post(f"/api/inspections/{inspection_id}/photos",
                          data={"section": "PACKING_ATTRIBUTE"},
                          files=[("files", (fname, fbytes, "image/jpeg"))])
        assert res.status_code == 201, res.text

    # 6. Buat Box Dimension Measurements & Upload foto berpasangan (Slot 1 & Slot 2)
    print("Mengunggah data Box Dimension beserta foto meteran...")
    box_axes = [("Length", 770), ("Width", 680), ("Height", 945)]
    box_photos = sample_photos["BOX_DIMENSION"]
    p_idx = 0
    for axis, val in box_axes:
        m_res = client.post(f"/api/inspections/{inspection_id}/measurements", json={
            "category": "BOX_DIMENSION",
            "label": axis,
            "value": val,
            "unit": "mm",
        })
        assert m_res.status_code == 201, m_res.text
        m_id = m_res.json()["id"]
        for slot in (1, 2):
            fname, fbytes = box_photos[p_idx]
            p_idx += 1
            pres = client.post(f"/api/inspections/{inspection_id}/photos",
                               data={"section": "BOX_DIMENSION", "measurement_id": m_id, "dimension_slot": slot},
                               files=[("files", (fname, fbytes, "image/jpeg"))])
            assert pres.status_code == 201, pres.text

    # 7. Upload Packing Method
    print("Mengunggah foto Packing Method...")
    for fname, fbytes in sample_photos["PACKING_METHOD"]:
        res = client.post(f"/api/inspections/{inspection_id}/photos",
                          data={"section": "PACKING_METHOD"},
                          files=[("files", (fname, fbytes, "image/jpeg"))])
        assert res.status_code == 201, res.text

    # 8. Upload Product View
    print("Mengunggah foto Product View...")
    for fname, fbytes in sample_photos["PRODUCT_VIEW"]:
        res = client.post(f"/api/inspections/{inspection_id}/photos",
                          data={"section": "PRODUCT_VIEW"},
                          files=[("files", (fname, fbytes, "image/jpeg"))])
        assert res.status_code == 201, res.text

    # 9. Upload Product Detail
    print("Mengunggah foto Product Detail...")
    for fname, fbytes in sample_photos["PRODUCT_DETAIL"]:
        res = client.post(f"/api/inspections/{inspection_id}/photos",
                          data={"section": "PRODUCT_DETAIL"},
                          files=[("files", (fname, fbytes, "image/jpeg"))])
        assert res.status_code == 201, res.text

    # 10. Upload Drain Hole
    print("Mengunggah foto Drain Hole...")
    for fname, fbytes in sample_photos["DRAIN_HOLE"]:
        res = client.post(f"/api/inspections/{inspection_id}/photos",
                          data={"section": "DRAIN_HOLE"},
                          files=[("files", (fname, fbytes, "image/jpeg"))])
        assert res.status_code == 201, res.text

    # 11. Upload Drawing
    print("Mengunggah foto Drawing...")
    fname, fbytes = sample_photos["DRAWING"][0]
    res = client.post(f"/api/inspections/{inspection_id}/photos",
                      data={"section": "DRAWING"},
                      files=[("files", (fname, fbytes, "image/jpeg"))])
    assert res.status_code == 201, res.text

    # 12. Buat Product Dimension Measurements & Upload foto berpasangan
    print("Mengunggah data Product Dimension beserta foto meteran...")
    prod_dims = [("Width:", 587), ("Depth:", 572), ("Height:", 764)]
    dim_photos = sample_photos["DIMENSION"]
    d_idx = 0
    for label, val in prod_dims:
        m_res = client.post(f"/api/inspections/{inspection_id}/measurements", json={
            "category": "DIMENSION",
            "label": label,
            "value": val,
            "unit": "mm",
        })
        assert m_res.status_code == 201, m_res.text
        m_id = m_res.json()["id"]
        for slot in (1, 2):
            fname, fbytes = dim_photos[d_idx]
            d_idx += 1
            pres = client.post(f"/api/inspections/{inspection_id}/photos",
                               data={"section": "DIMENSION", "measurement_id": m_id, "dimension_slot": slot},
                               files=[("files", (fname, fbytes, "image/jpeg"))])
            assert pres.status_code == 201, pres.text

    # 13. Level Check
    print("Mengunggah foto Level Check...")
    for fname, fbytes in sample_photos["LEVEL_CHECK"]:
        res = client.post(f"/api/inspections/{inspection_id}/photos",
                          data={"section": "LEVEL_CHECK"},
                          files=[("files", (fname, fbytes, "image/jpeg"))])
        assert res.status_code == 201, res.text

    # 14. Mal Check
    print("Mengunggah foto Mal Check...")
    for fname, fbytes in sample_photos["MAL_CHECK"]:
        res = client.post(f"/api/inspections/{inspection_id}/photos",
                          data={"section": "MAL_CHECK"},
                          files=[("files", (fname, fbytes, "image/jpeg"))])
        assert res.status_code == 201, res.text

    # 15. Radius Mal (R44 & R10)
    print("Mengunggah foto Radius Check...")
    for fname, fbytes in sample_photos["RADIUS"]:
        res = client.post(f"/api/inspections/{inspection_id}/photos",
                          data={"section": "RADIUS"},
                          files=[("files", (fname, fbytes, "image/jpeg"))])
        assert res.status_code == 201, res.text

    # 16. Tambahkan Issue (Cacat / Temuan Masalah)
    print("Mencatat temuan cacat (Issue)...")
    # Issue 1: Rough
    iss1 = client.post(f"/api/inspections/{inspection_id}/issues", json={
        "defect_type": "Rough",
        "quantity": 2,
        "cause": "Uneven powder coating spray gun distance / minor dust on conveyor line",
        "repair": "Light fine sanding (grit 800) and spot touch-up re-coating",
        "cap": "Improve spray booth air filtration and enforce 100% surface air blowing prior to powder application",
    })
    assert iss1.status_code == 201, iss1.text
    iss1_id = iss1.json()["id"]
    rfname, rfbytes = sample_photos["ISSUE"][0]
    p_iss1 = client.post(f"/api/inspections/{inspection_id}/photos",
                         data={"section": "ISSUE", "issue_id": iss1_id},
                         files=[("files", (rfname, rfbytes, "image/jpeg"))])
    assert p_iss1.status_code == 201, p_iss1.text

    # Issue 2: Scratch
    iss2 = client.post(f"/api/inspections/{inspection_id}/issues", json={
        "defect_type": "Scratch",
        "quantity": 1,
        "cause": "Minor friction during transit on factory transfer trolley",
        "repair": "Polishing compound and touch up protective lacquer",
        "cap": "Install rubber protector padding on all factory transport carts",
    })
    assert iss2.status_code == 201, iss2.text
    iss2_id = iss2.json()["id"]
    sfname, sfbytes = sample_photos["ISSUE"][1]
    p_iss2 = client.post(f"/api/inspections/{inspection_id}/photos",
                         data={"section": "ISSUE", "issue_id": iss2_id},
                         files=[("files", (sfname, sfbytes, "image/jpeg"))])
    assert p_iss2.status_code == 201, p_iss2.text

    # 17. Validasi Preflight
    print("Menjalankan validasi pra-cetak (preflight review)...")
    val_res = client.post(f"/api/inspections/{inspection_id}/validate")
    assert val_res.status_code == 200, val_res.text
    val_data = val_res.json()
    print("Hasil Validasi Pra-cetak:")
    print("Errors:", val_data["errors"])
    print("Warnings:", val_data["warnings"])
    assert val_data["errors"] == [], f"Ada error validasi: {val_data['errors']}"

    # 18. Generate Dokumen Word (.docx)
    print("Mencetak dokumen laporan Word resmi dari aplikasi...")
    rep_res = client.post(f"/api/inspections/{inspection_id}/report/docx")
    assert rep_res.status_code == 200, rep_res.text
    docx_bytes = rep_res.content

    # 19. Simpan file Word ke folder Downloads pengguna
    downloads_dir = Path(os.environ.get("USERPROFILE", str(Path.home()))) / "Downloads"
    downloads_dir.mkdir(parents=True, exist_ok=True)
    out_docx_path = downloads_dir / "Contoh_Laporan_QC_Gabungan_Inspectra.docx"
    out_docx_path.write_bytes(docx_bytes)
    print(f"File laporan Word berhasil disimpan di: {out_docx_path}")
    print(f"Ukuran file: {len(docx_bytes)} bytes")

    # 20. Verifikasi isi dokumen
    doc = Document(io.BytesIO(docx_bytes))
    print(f"Jumlah seksi dokumen: {len(doc.sections)}")
    print(f"Jumlah tabel: {len(doc.tables)}")
    print(f"Jumlah paragraf: {len(doc.paragraphs)}")
    print("CONTOH LAPORAN SELESAI DIBUAT DENGAN SUKSES!")


if __name__ == "__main__":
    main()
