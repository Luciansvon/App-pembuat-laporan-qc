import hashlib
import io
import sys
from pathlib import Path
import httpx
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

BASE_URL = "http://127.0.0.1:8000"

def make_sample_image(text="SAMPLE"):
    img = Image.new("RGB", (800, 600), color=(183, 166, 142))
    buf = io.BytesIO()
    img.save(buf, format="JPEG")
    return buf.getvalue()

def main():
    client = httpx.Client(base_url=BASE_URL, timeout=30.0)
    
    # 1. Health check
    res = client.get("/api/health")
    print(f"Health check: {res.status_code} {res.json()}")
    assert res.status_code == 200

    # 2. Buat Customer
    customer_res = client.post("/api/customers", json={"code": "POLIFORM", "name": "Poliform QC"})
    if customer_res.status_code == 409:
        # maybe exists, get it
        customers = client.get("/api/customers").json()
        cust = next(c for c in customers if c["code"] == "POLIFORM")
        cust_id = cust["id"]
        print(f"Customer existing: {cust['name']} ({cust_id})")
    else:
        assert customer_res.status_code == 201, customer_res.text
        cust_id = customer_res.json()["id"]
        print(f"Customer created: {cust_id}")

    # 3. Buat Produk
    prod_res = client.post("/api/products", json={
        "customer_id": cust_id,
        "name": "Kursi Santai Wood",
        "standard_dimensions": {"L": 850, "W": 750, "H": 900}
    })
    assert prod_res.status_code == 201, prod_res.text
    prod_id = prod_res.json()["id"]
    print(f"Product created: {prod_id}")

    # 4. Buat Inspeksi
    insp_res = client.post("/api/inspections", json={
        "product_id": prod_id,
        "po": "PO-2026-001",
        "inspection_type": "FINAL",
        "date": "2026-10-12",
        "location": "Workshop Bima",
        "qc_name": "Bima",
        "quantity": 10,
        "aql": "1.5",
        "box_dimensions": {"L": 900, "W": 800, "H": 950}
    })
    assert insp_res.status_code == 201, insp_res.text
    insp = insp_res.json()
    insp_id = insp["id"]
    print(f"Inspection created: {insp['inspection_number']} ({insp_id})")

    # 5. Tab Photos: isi 2 foto Product View
    img_data = make_sample_image("Product View")
    photo_res = client.post(
        f"/api/inspections/{insp_id}/photos",
        data={"section": "PRODUCT_VIEW"},
        files=[
            ("files", ("photo1.jpg", img_data, "image/jpeg")),
            ("files", ("photo2.jpg", img_data, "image/jpeg"))
        ]
    )
    assert photo_res.status_code == 201, photo_res.text
    print(f"Uploaded 2 Product View photos: {len(photo_res.json())} photos saved.")

    # 6. Tab Measurement: tambah ukuran Length + isi 2 slot fotonya
    meas_res = client.post(f"/api/inspections/{insp_id}/measurements", json={
        "category": "DIMENSION",
        "label": "Length",
        "unit": "mm",
        "standard_value": 850.0,
        "tolerance_plus": 5.0,
        "tolerance_minus": 5.0,
        "value": 852.0,
        "notes": "Real dimension length check"
    })
    assert meas_res.status_code == 201, meas_res.text
    meas_id = meas_res.json()["id"]
    print(f"Added measurement Length: {meas_id}, result: {meas_res.json()['result']}")

    # Isi 2 slot foto measurement
    dim_photo_1 = client.post(
        f"/api/inspections/{insp_id}/photos",
        data={"section": "DIMENSION", "measurement_id": meas_id, "dimension_slot": "1"},
        files=[("files", ("dim1.jpg", img_data, "image/jpeg"))]
    )
    assert dim_photo_1.status_code == 201, dim_photo_1.text

    dim_photo_2 = client.post(
        f"/api/inspections/{insp_id}/photos",
        data={"section": "DIMENSION", "measurement_id": meas_id, "dimension_slot": "2"},
        files=[("files", ("dim2.jpg", img_data, "image/jpeg"))]
    )
    assert dim_photo_2.status_code == 201, dim_photo_2.text
    print("Uploaded 2 slot photos for measurement Length.")

    # 7. Review & Validate
    val_res = client.post(f"/api/inspections/{insp_id}/validate")
    assert val_res.status_code == 200, val_res.text
    val_data = val_res.json()
    print(f"Validation errors: {val_data['errors']}")
    print(f"Validation warnings: {val_data['warnings']}")
    assert len(val_data["errors"]) == 0, f"Validation failed with errors: {val_data['errors']}"

    # 8. Unduh DOCX
    docx_res = client.post(f"/api/inspections/{insp_id}/report/docx")
    assert docx_res.status_code == 200, docx_res.text
    
    out_path = Path("C:/Users/shint/AppData/Local/Temp/opencode/live_qa_report.docx")
    out_path.write_bytes(docx_res.content)
    print(f"Downloaded DOCX to {out_path} ({len(docx_res.content)} bytes).")

    # 9. Verifikasi isi DOCX
    doc = Document(io.BytesIO(docx_res.content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    print(f"Paragraphs in DOCX: {paragraphs}")

    # Cek heading Product Dimension
    dim_heading = next((p for p in doc.paragraphs if p.text == "Product Dimension"), None)
    assert dim_heading is not None, "Product Dimension heading not found!"
    print(f"Product Dimension page_break_before: {dim_heading.paragraph_format.page_break_before}")
    assert dim_heading.paragraph_format.page_break_before is True, "Product Dimension MUST have page_break_before=True!"

    # Cek header tabel DIM/BOX
    header_table = doc.sections[0].header.tables[0]
    header_cells = [cell.text.strip() for row in header_table.rows for cell in row.cells]
    print(f"Header cells: {header_cells}")
    dim_found = any("DIM" in c for c in header_cells)
    box_found = any("BOX" in c for c in header_cells)
    print(f"DIM found in header: {dim_found}, BOX found in header: {box_found}")
    assert dim_found, "DIM row not found in header table"
    assert box_found, "BOX row not found in header table"

    print("\n>>> SEMUA UJI ALUR & VERIFIKASI DOCX LOLOS DENGAN SUKSES! <<<")

if __name__ == "__main__":
    main()
