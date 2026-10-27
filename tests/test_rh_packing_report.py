from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from app.core.config import Settings
from app.main import create_app


def image_bytes(width: int = 400, height: int = 300, color: str = "#8B5A2B") -> bytes:
    image = Image.new("RGB", (width, height), color=color)
    stream = io.BytesIO()
    image.save(stream, format="JPEG")
    return stream.getvalue()


def test_rh_packing_report_generation(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path)
    app = create_app(settings)
    with TestClient(app) as client:
        # Create RH customer & product
        customer_resp = client.post("/api/customers", json={"code": "RH", "name": "Restoration Hardware"})
        assert customer_resp.status_code == 201, customer_resp.text
        customer_id = customer_resp.json()["id"]

        product_resp = client.post("/api/products", json={
            "customer_id": customer_id,
            "name": "Marbella Side Chair",
            "standard_dimensions": {"W": 840, "D": 683, "H": 1003}
        })
        assert product_resp.status_code == 201, product_resp.text
        product_id = product_resp.json()["id"]

        # Create packing inspection
        insp_resp = client.post("/api/inspections", json={
            "product_id": product_id,
            "po": "5273162",
            "inspection_type": "PACKING",
            "report_type": "PACKING",
            "date": "2026-08-04",
            "location": "QSF",
            "qc_name": "Ramanov",
            "quantity": 47,
            "aql": None,
            "box_dimensions": {"W": 840, "D": 683, "H": 1003}
        })
        assert insp_resp.status_code == 201, insp_resp.text
        insp = insp_resp.json()
        item_id = insp["id"]

        # Add 3 Box Dimension measurements (Width, Depth, Height) with 2 photos each
        for axis, val in [("Width", 840), ("Depth", 683), ("Height", 1003)]:
            m_resp = client.post(f"/api/inspections/{item_id}/measurements", json={
                "category": "BOX_DIMENSION", "label": axis, "value": val, "unit": "mm"
            })
            assert m_resp.status_code == 201, m_resp.text
            m_id = m_resp.json()["id"]

            photos_resp = client.post(f"/api/inspections/{item_id}/photos",
                data={"section": "BOX_DIMENSION", "measurement_id": m_id},
                files=[("files", (f"{axis}-{slot}.jpg", image_bytes(), "image/jpeg")) for slot in (1, 2)])
            assert photos_resp.status_code == 201, photos_resp.text

        # Add photos for PACKING_ATTRIBUTE (Packing View, Label, Certification)
        attr_photos = client.post(f"/api/inspections/{item_id}/photos",
            data={"section": "PACKING_ATTRIBUTE"},
            files=[("files", (f"attr-{i}.jpg", image_bytes(), "image/jpeg")) for i in range(4)])
        assert attr_photos.status_code == 201, attr_photos.text

        # Add photos for PACKING_METHOD
        method_photos = client.post(f"/api/inspections/{item_id}/photos",
            data={"section": "PACKING_METHOD"},
            files=[("files", (f"method-{i}.jpg", image_bytes(), "image/jpeg")) for i in range(6)])
        assert method_photos.status_code == 201, method_photos.text

        # Add photos for INNER_SKU_GLIDE
        inner_photos = client.post(f"/api/inspections/{item_id}/photos",
            data={"section": "INNER_SKU_GLIDE"},
            files=[("files", (f"inner-{i}.jpg", image_bytes(), "image/jpeg")) for i in range(6)])
        assert inner_photos.status_code == 201, inner_photos.text

        # An MC measurement must not re-enable the table disabled by rh_packing config
        mc_resp = client.post(f"/api/inspections/{item_id}/measurements", json={
            "category": "MC", "label": "Moisture Content", "value": 10.5, "unit": "%"
        })
        assert mc_resp.status_code == 201, mc_resp.text

        # Validate
        val_resp = client.post(f"/api/inspections/{item_id}/validate")
        assert val_resp.status_code == 200, val_resp.text
        val_data = val_resp.json()
        assert val_data["errors"] == []

        # Generate report
        report_resp = client.post(f"/api/inspections/{item_id}/report/docx")
        assert report_resp.status_code == 200, report_resp.text
        assert report_resp.content.startswith(b"PK")

        word = Document(io.BytesIO(report_resp.content))
        header_table = word.sections[0].header.tables[0]
        header_text = " ".join(cell.text for row in header_table.rows for cell in row.cells)

        assert "{{" not in header_text
        assert "Marbella Side Chair" in header_text
        assert "PACKING" in header_text
        assert "5273162" in header_text
        assert "260804-QSF" in header_text
        assert "Ramanov" in header_text
        assert "840" in header_text
        assert "683" in header_text
        assert "1003" in header_text

        # Check body headings and structure
        headings = [p.text for p in word.paragraphs if p.text.strip()]
        assert "ATRIBUTE (Box & Pack)" in headings
        assert "Box Dimension" in headings
        assert "Packing Method" in headings
        assert "Inner SKU & Fixed Glide" in headings
        assert "Conclusion" not in headings
        assert "Measurement and Test Results" not in headings

        # Check total tables in body (photo grids + dimension blocks)
        assert len(word.tables) >= 5
