"""Test for RODA combined inspection and packing report."""
from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from app.core.config import Settings
from app.main import create_app


def synthetic_image_bytes(color: str = "#264653") -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (64, 48), color).save(output, format="JPEG")
    return output.getvalue()


def test_roda_combined_report_generation(tmp_path: Path) -> None:
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        # 1. Create RODA customer & product
        customer_res = client.post("/api/customers", json={"code": "RODA", "name": "RODA"})
        assert customer_res.status_code == 201
        customer_id = customer_res.json()["id"]

        product_res = client.post("/api/products", json={
            "customer_id": customer_id,
            "name": "ERICE STACKING CHAIR",
            "standard_dimensions": {"W": 587, "D": 572, "H": 764},
        })
        assert product_res.status_code == 201
        product_id = product_res.json()["id"]

        # 2. Create inspection with RODA metadata matching the reference document
        insp_res = client.post("/api/inspections", json={
            "product_id": product_id,
            "po": "248",
            "inspection_type": "FINAL",
            "report_type": "COMBINED",
            "date": "2024-04-26",
            "location": "QSF",
            "qc_name": "GENTA",
            "quantity": 200,
            "product_dimensions": {"W": 587, "D": 572, "H": 764},
            "box_dimensions": {"L": 770, "W": 680, "H": 945},
        })
        assert insp_res.status_code == 201
        inspection_id = insp_res.json()["id"]

        # 3. Add Packing Attribute photos (6 photos)
        for i in range(6):
            res = client.post(f"/api/inspections/{inspection_id}/photos",
                              data={"section": "PACKING_ATTRIBUTE"},
                              files=[("files", (f"attr_{i}.jpg", synthetic_image_bytes("#2a9d8f"), "image/jpeg"))])
            assert res.status_code == 201

        # 4. Add Box Dimension measurements (Length, Width, Height) with 2 photos each
        for axis, val in [("Length", 770), ("Width", 680), ("Height", 945)]:
            m_res = client.post(f"/api/inspections/{inspection_id}/measurements", json={
                "category": "BOX_DIMENSION",
                "label": axis,
                "value": val,
                "unit": "mm",
            })
            assert m_res.status_code == 201
            m_id = m_res.json()["id"]
            for slot in (1, 2):
                p_res = client.post(f"/api/inspections/{inspection_id}/photos",
                                    data={"section": "BOX_DIMENSION", "measurement_id": m_id, "dimension_slot": slot},
                                    files=[("files", (f"box_{axis}_{slot}.jpg", synthetic_image_bytes("#e76f51"), "image/jpeg"))])
                assert p_res.status_code == 201

        # 5. Add Packing Method photos (6 photos)
        for i in range(6):
            res = client.post(f"/api/inspections/{inspection_id}/photos",
                              data={"section": "PACKING_METHOD"},
                              files=[("files", (f"pack_{i}.jpg", synthetic_image_bytes("#f4a261"), "image/jpeg"))])
            assert res.status_code == 201

        # 6. Add Product View photos
        for i in range(4):
            res = client.post(f"/api/inspections/{inspection_id}/photos",
                              data={"section": "PRODUCT_VIEW"},
                              files=[("files", (f"prod_{i}.jpg", synthetic_image_bytes("#457b9d"), "image/jpeg"))])
            assert res.status_code == 201

        # 7. Add Product Dimension measurements with 2 photos each
        for label, val in [("Width:", 587), ("Depth:", 572), ("Height:", 764)]:
            m_res = client.post(f"/api/inspections/{inspection_id}/measurements", json={
                "category": "DIMENSION",
                "label": label,
                "value": val,
                "unit": "mm",
            })
            assert m_res.status_code == 201
            m_id = m_res.json()["id"]
            for slot in (1, 2):
                p_res = client.post(f"/api/inspections/{inspection_id}/photos",
                                    data={"section": "DIMENSION", "measurement_id": m_id, "dimension_slot": slot},
                                    files=[("files", (f"dim_{slot}.jpg", synthetic_image_bytes("#1d3557"), "image/jpeg"))])
                assert p_res.status_code == 201

        # 8. Add an Issue (Rough) with photo, cause, repair, CAP
        issue_res = client.post(f"/api/inspections/{inspection_id}/issues", json={
            "defect_type": "Rough",
            "cause": "Uneven powder coating application/Un clean surface",
            "repair": "Sanding and re-coating",
            "cap": "Make sure to improve powder coating control.",
        })
        assert issue_res.status_code == 201
        issue_id = issue_res.json()["id"]
        res = client.post(f"/api/inspections/{inspection_id}/photos",
                          data={"section": "ISSUE", "issue_id": issue_id},
                          files=[("files", ("issue_rough.jpg", synthetic_image_bytes("#e63946"), "image/jpeg"))])
        assert res.status_code == 201

        # 9. Validate inspection preflight
        review_res = client.post(f"/api/inspections/{inspection_id}/validate")
        assert review_res.status_code == 200, review_res.text
        review_data = review_res.json()
        assert review_data["errors"] == [], f"Validation errors: {review_data['errors']}"

        # 10. Generate DOCX report
        report_res = client.post(f"/api/inspections/{inspection_id}/report/docx")
        assert report_res.status_code == 200, report_res.text
        doc = Document(io.BytesIO(report_res.content))

        # Check Header
        header_table = doc.sections[0].header.tables[0]
        header_text = [c.text.strip().replace("\n", " ") for row in header_table.rows for c in row.cells]
        full_header = " | ".join(header_text)
        assert "ERICE STACKING CHAIR" in full_header
        assert "RODA" in full_header
        assert "248" in full_header
        assert "260424-QSF" in full_header
        assert "GENTA" in full_header
        assert "200" in full_header
        # Sumbu W D H dan L W H
        assert "587" in full_header
        assert "572" in full_header
        assert "764" in full_header
        assert "770" in full_header
        assert "680" in full_header
        assert "945" in full_header

        # Check Body Paragraphs & Headings
        body_text = "\n".join(p.text for p in doc.paragraphs)
        assert "ATRIBUTE( Box, packing ,Label, Position Logo)" in body_text
        assert "Box Dimension" in body_text
        assert "Packing Method" in body_text
        assert "Product View and Detail" in body_text or "Product View" in body_text
        assert "Product Dimension" in body_text
        assert "Rough" in "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)


def test_any_customer_can_use_combined_report(tmp_path: Path) -> None:
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        # Customer baru apapun (misalnya BIMA-JAYA)
        c_res = client.post("/api/customers", json={"code": "BIMA-JAYA", "name": "Bima Jaya Furniture"})
        assert c_res.status_code == 201
        p_res = client.post("/api/products", json={"customer_id": c_res.json()["id"], "name": "Custom Dining Table"})
        assert p_res.status_code == 201
        
        # Buat inspeksi dengan report_type="COMBINED"
        insp_res = client.post("/api/inspections", json={
            "product_id": p_res.json()["id"],
            "po": "PO-999",
            "inspection_type": "FINAL",
            "report_type": "COMBINED",
            "date": "2026-10-21",
            "qc_name": "Bima QC",
            "quantity": 50,
        })
        assert insp_res.status_code == 201
        insp_id = insp_res.json()["id"]
        assert insp_res.json()["report_type"] == "COMBINED"

        # Upload photo in PACKING_ATTRIBUTE
        p_upload = client.post(f"/api/inspections/{insp_id}/photos",
                               data={"section": "PACKING_ATTRIBUTE"},
                               files=[("files", ("attr.jpg", synthetic_image_bytes(), "image/jpeg"))])
        assert p_upload.status_code == 201

        # Preflight validation passes without error
        rev = client.post(f"/api/inspections/{insp_id}/validate")
        assert rev.status_code == 200
        assert rev.json()["errors"] == []

