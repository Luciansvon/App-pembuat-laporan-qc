"""Test MC range caption and inspected product photo rendering in reports."""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from app.core.config import Settings
from app.main import create_app


def synthetic_image_bytes(color: str = "#3a5a40") -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (64, 48), color).save(output, format="JPEG")
    return output.getvalue()


def test_mc_range_caption_and_inspected_product_rendering(tmp_path: Path) -> None:
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        # 1. Customer & Product
        cust = client.post("/api/customers", json={"code": "UNOPIU", "name": "UNOPIU"}).json()
        prod = client.post("/api/products", json={
            "customer_id": cust["id"],
            "name": "ELBA 80 SQUARE TABLE",
            "standard_dimensions": {"W": 800, "D": 800, "H": 760},
        }).json()

        # 2. Inspection with inspected_quantity = 7
        insp = client.post("/api/inspections", json={
            "product_id": prod["id"],
            "po": "NOVEMBER",
            "inspection_type": "INLINE",
            "date": "2026-10-22",
            "qc_name": "Ramanov",
            "quantity": 65,
            "inspected_quantity": 7,
            "aql": "7",
            "product_dimensions": {"W": 800, "D": 800, "H": 760},
        }).json()
        insp_id = insp["id"]

        # 3. Add 2 MC photos with captions 8.5% and 11.4%
        p1 = client.post(f"/api/inspections/{insp_id}/photos",
                         data={"section": "MC"},
                         files=[("files", ("mc1.jpg", synthetic_image_bytes("#a3b18a"), "image/jpeg"))]).json()[0]
        client.patch(f"/api/photos/{p1['id']}", json={"caption": "8.5%"})

        p2 = client.post(f"/api/inspections/{insp_id}/photos",
                         data={"section": "MC"},
                         files=[("files", ("mc2.jpg", synthetic_image_bytes("#588157"), "image/jpeg"))]).json()[0]
        client.patch(f"/api/photos/{p2['id']}", json={"caption": "11.4%"})

        # 4. Add INSPECTED_PRODUCT photo with default or custom caption
        p3 = client.post(f"/api/inspections/{insp_id}/photos",
                         data={"section": "INSPECTED_PRODUCT"},
                         files=[("files", ("inspect_sample.jpg", synthetic_image_bytes("#344e41"), "image/jpeg"))]).json()[0]

        # 5. Preflight validation
        rev = client.post(f"/api/inspections/{insp_id}/validate")
        assert rev.status_code == 200
        assert rev.json()["errors"] == []

        # 6. Generate DOCX
        rep = client.post(f"/api/inspections/{insp_id}/report/docx")
        assert rep.status_code == 200, rep.text
        doc = Document(io.BytesIO(rep.content))

        # Check body paragraphs
        body_paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Verify MC Range caption is generated properly
        assert any("Range MC Check is 8.5% - 11.4%" in p for p in body_paras), f"MC range not found: {body_paras}"

        # Verify no Measurement and Test Results table
        assert not any("Measurement and Test Results" in p for p in body_paras), "Measurement table should not be present"

        # Check table captions (inspect 7 pcs)
        table_texts = [cell.text.strip() for t in doc.tables for row in t.rows for cell in row.cells if cell.text.strip()]
        assert any("inspect 7 pcs" in t for t in table_texts), f"inspect 7 pcs not found in tables: {table_texts}"
