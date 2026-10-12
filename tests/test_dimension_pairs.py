from __future__ import annotations

import hashlib
import io
from pathlib import Path
import sqlite3

from alembic import command
from alembic.config import Config
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from app.core.config import Settings
from app.main import create_app
from test_workflow import create_basic_inspection, image_bytes


def upload(client: TestClient, inspection: str, measurement: str, *, slot: int | None = None,
           count: int = 1, color: str = '#72584b'):
    data = {'section': 'DIMENSION', 'measurement_id': measurement}
    if slot is not None:
        data['dimension_slot'] = str(slot)
    return client.post(f'/api/inspections/{inspection}/photos', data=data,
        files=[('files', (f'photo-{i}.jpg', image_bytes(color), 'image/jpeg')) for i in range(count)])


def test_each_dimension_has_its_own_pair_even_when_upload_order_is_mixed(tmp_path: Path) -> None:
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        inspection = create_basic_inspection(client)['id']
        measurements = [client.post(f'/api/inspections/{inspection}/measurements', json={
            'category': 'DIMENSION', 'label': label, 'unit': 'mm', 'sort_order': index,
            'value': [1200, 600, 750, 450][index],
        }).json() for index, label in enumerate(['Length', 'Width', 'Depth', 'Height'])]
        expected: dict[tuple[int, int], bytes] = {}
        for index, slot in [(3, 2), (0, 1), (1, 2), (2, 1), (0, 2), (3, 1), (2, 2), (1, 1)]:
            response = upload(client, inspection, measurements[index]['id'], slot=slot, color=f'#{index + 1}0{slot}0a0')
            assert response.status_code == 201, response.text
            photo = response.json()[0]
            assert photo['dimension_slot'] == slot
            expected[index, slot] = client.get(f"/api/photos/{photo['id']}/file").content
            assert client.patch(f"/api/photos/{photo['id']}", json={'sort_order': 100 - index * 10 - slot}).status_code == 200
        assert client.post(f'/api/inspections/{inspection}/validate').json()['errors'] == []
        response = client.post(f'/api/inspections/{inspection}/report/docx')
        assert response.status_code == 200, response.text
        word = Document(io.BytesIO(response.content))
        assert len(word.inline_shapes) == 8
        for index, measurement in enumerate(measurements):
            table = next(table for table in word.tables if table.cell(0, 0).text.startswith(measurement['label']))
            assert table.cell(0, 0)._tc is table.cell(0, 1)._tc
            for slot in (1, 2):
                cell = table.cell(1, slot - 1)
                pictures = cell._tc.xpath('.//a:blip')
                assert len(pictures) == 1
                relationship = pictures[0].get(qn('r:embed'))
                assert word.part.related_parts[relationship].blob == expected[index, slot]
        assert hashlib.sha256(client.post(f'/api/inspections/{inspection}/report/docx').content).digest() == hashlib.sha256(response.content).digest()


def test_dimension_requires_two_and_rejects_extra_upload_atomically(tmp_path: Path) -> None:
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        inspection = create_basic_inspection(client)['id']
        measurement = client.post(f'/api/inspections/{inspection}/measurements', json={'category': 'DIMENSION', 'label': 'Length'}).json()
        assert upload(client, inspection, measurement['id']).status_code == 201
        assert 'exactly 2 photos' in str(client.post(f'/api/inspections/{inspection}/validate').json()['errors'])
        assert client.post(f'/api/inspections/{inspection}/report/docx').status_code == 409
        assert upload(client, inspection, measurement['id'], count=2).status_code == 409
        assert len(client.get(f'/api/inspections/{inspection}').json()['photos']) == 1
        assert len(list(tmp_path.rglob('photos/*.jpg'))) == 1
        assert upload(client, inspection, measurement['id'], slot=2).status_code == 201
        assert upload(client, inspection, measurement['id']).status_code == 409
        assert client.post(f'/api/inspections/{inspection}/report/docx').status_code == 200
        other = create_basic_inspection(client, 'OTHER-QA')['id']
        assert upload(client, other, measurement['id']).status_code == 422
        mc = client.post(f'/api/inspections/{inspection}/measurements', json={'category': 'MC', 'label': 'MC'}).json()
        assert upload(client, inspection, mc['id']).status_code == 422


def test_dimension_section_starts_on_new_page_after_product_view(tmp_path: Path) -> None:
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        inspection = create_basic_inspection(client, 'RH')['id']
        views = client.post(f'/api/inspections/{inspection}/photos', data={'section': 'PRODUCT_VIEW'},
            files=[('files', (f'view-{i}.jpg', image_bytes(), 'image/jpeg')) for i in range(2)])
        assert views.status_code == 201, views.text
        measurement = client.post(f'/api/inspections/{inspection}/measurements', json={
            'category': 'DIMENSION', 'label': 'Length', 'unit': 'mm', 'value': 1965}).json()
        assert upload(client, inspection, measurement['id'], count=2).status_code == 201
        assert client.post(f'/api/inspections/{inspection}/validate').json()['errors'] == []
        response = client.post(f'/api/inspections/{inspection}/report/docx')
        assert response.status_code == 200, response.text
        word = Document(io.BytesIO(response.content))
        view_heading = next(p for p in word.paragraphs if p.text == 'Product View')
        dimension_heading = next(p for p in word.paragraphs if p.text == 'Product Dimension')
        assert view_heading.paragraph_format.page_break_before in (None, False)
        assert dimension_heading.paragraph_format.page_break_before is True


def test_old_photos_survive_migration_and_explicit_assignment(tmp_path: Path, monkeypatch) -> None:
    settings = Settings(data_dir=tmp_path)
    with TestClient(create_app(settings)) as client:
        inspection = create_basic_inspection(client)['id']
        photo = client.post(f'/api/inspections/{inspection}/photos', data={'section': 'PRODUCT_VIEW'},
            files=[('files', ('old.jpg', image_bytes(), 'image/jpeg'))]).json()[0]
        measurement = client.post(f'/api/inspections/{inspection}/measurements', json={'category': 'MC', 'label': 'Old MC'}).json()
        reading = client.post(f"/api/measurements/{measurement['id']}/readings", json={'value': 10, 'photo_id': photo['id']}).json()
        before = client.get(f"/api/photos/{photo['id']}/file").content
    monkeypatch.setenv('QC_DATA_DIR', str(tmp_path))
    root = Path(__file__).resolve().parents[1]
    config = Config(str(root / 'backend/alembic.ini'))
    config.set_main_option('script_location', str(root / 'backend/migrations'))
    command.downgrade(config, '446c624e36d9')
    with sqlite3.connect(settings.database_path) as database:
        database.execute("UPDATE photos SET section='DIMENSION' WHERE id=?", (photo['id'],))
    with TestClient(create_app(settings)) as client:
        assert client.get(f"/api/photos/{photo['id']}/file").content == before
        assert client.post(f'/api/inspections/{inspection}/report/docx').status_code == 409
        dimension = client.post(f'/api/inspections/{inspection}/measurements', json={'category': 'DIMENSION', 'label': 'Depth'}).json()
        assert client.patch(f"/api/photos/{photo['id']}", json={'measurement_id': dimension['id']}).status_code == 200
        assert upload(client, inspection, dimension['id']).status_code == 201
        assert client.delete(f"/api/measurements/{dimension['id']}", headers={'X-Confirm-Delete': 'true'}).status_code == 200
        retained = client.get(f'/api/inspections/{inspection}').json()
        assert all(photo['section'] == 'PRODUCT_VIEW' and photo['measurement_id'] is None for photo in retained['photos'])
        assert retained['measurements'][0]['readings'][0]['id'] == reading['id']
        assert client.get(f"/api/photos/{photo['id']}/file").content == before
