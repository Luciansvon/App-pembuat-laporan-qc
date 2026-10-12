"""Generated photo bounds must fit the containing cells for every supported count."""
import io
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from fastapi.testclient import TestClient
from PIL import Image

from app.core.config import Settings
from app.main import create_app
from app.models.tables import InspectionPhoto, Issue
from app.report.generator import PHOTO_PADDING_CM, issue_block, photo_grid, row_spans
from test_workflow import create_basic_inspection, image_bytes


@pytest.mark.parametrize("count", range(1, 7))
@pytest.mark.parametrize("kind", ["issue", "section"])
def test_photo_bounds_and_odd_final_row(tmp_path: Path, count: int, kind: str) -> None:
    shapes = [(6000, 1000), (1000, 6000), (1600, 1600), (3000, 2000), (6000, 1000), (1200, 2400)]
    photos = []
    for index, size in enumerate(shapes[:count]):
        path = tmp_path / f"{index}.jpg"
        Image.new("RGB", size, (30 + index * 25, 120, 90)).save(path)
        photos.append(InspectionPhoto(inspection_id="QA", section="ISSUE", file_path=path.name,
                                      caption=f"Caption {index + 1}"))
    document = Document()
    settings = Settings(data_dir=tmp_path)
    if kind == "issue":
        issue_block(document, Issue(inspection_id="QA", defect_type="QA", cause="CAUSE",
                                    repair="REPAIR", cap="CAP"), photos, settings)
        grid = document.tables[0].rows[1].cells[0].tables[0]
        # The paragraph between the label row and nested grid must remain present.
        cell = document.tables[0].rows[1].cells[0]._tc
        assert cell[1].tag == qn("w:p")
    else:
        photo_grid(document, photos, settings)
        grid = document.tables[0]
    path = tmp_path / "report.docx"
    document.save(path)
    reloaded = Document(path)
    assert len(reloaded.inline_shapes) == count
    assert not reloaded.element.xpath(".//wp:anchor")
    for cell in reloaded.element.xpath(".//w:tc"):
        width = cell.tcPr.tcW.w
        for inline in cell.xpath("./w:p//wp:inline"):
            extent = inline.xpath("./wp:extent")[0]
            available = width * 635 - Cm(2 * PHOTO_PADDING_CM).emu
            assert int(extent.get("cx")) <= available + 635
    for shape, (width, height) in zip(reloaded.inline_shapes, shapes, strict=False):
        assert abs(shape.width / shape.height - width / height) < 0.001
    photo_rows = [row for row in grid.rows if row._tr.xpath(".//wp:inline")]
    assert len(photo_rows) == (count + 1) // 2
    if count % 2:
        assert photo_rows[-1].cells[0]._tc is photo_rows[-1].cells[1]._tc
        # A wide final photo must use the merged cell, not a half-width limit.
        if count == 5:
            assert reloaded.inline_shapes[-1].width > Cm(12)


@pytest.mark.parametrize("cols,count,expected_spans", [
    (2, 2, [1, 1]), (3, 2, [2, 1]), (4, 2, [2, 2]),
    (5, 3, [2, 2, 1]), (8, 1, [8]), (8, 8, [1] * 8),
])
def test_row_spans_share_columns_evenly(cols: int, count: int, expected_spans: list[int]) -> None:
    assert row_spans(cols, count) == expected_spans
    assert sum(expected_spans) == cols


@pytest.mark.parametrize("cols", [1, 3, 4, 8])
def test_layout_cols_render_requested_columns(tmp_path: Path, cols: int) -> None:
    count = 7
    photos = []
    for index in range(count):
        path = tmp_path / f"{index}.jpg"
        Image.new("RGB", (1200, 900), (40 + index * 20, 120, 90)).save(path)
        photos.append(InspectionPhoto(inspection_id="QA", section="PRODUCT_VIEW", file_path=path.name,
                                      caption=f"Caption {index + 1}"))
    document = Document()
    photo_grid(document, photos, Settings(data_dir=tmp_path), cols=cols)
    grid = document.tables[0]
    assert len(grid.columns) == cols
    photo_rows = [row for row in grid.rows if row._tr.xpath(".//wp:inline")]
    assert len(photo_rows) == (count + cols - 1) // cols
    assert len(document.inline_shapes) == count
    caption_rows = [row for row in grid.rows if not row._tr.xpath(".//wp:inline")]
    assert len(caption_rows) == len(photo_rows)
    for caption_row in caption_rows:
        assert caption_row.cells[0]._tc is caption_row.cells[cols - 1]._tc
        assert caption_row.cells[0].text.strip() != ""


def test_layout_partial_row_merges_cells(tmp_path: Path) -> None:
    photos = []
    for index in range(3):
        path = tmp_path / f"{index}.jpg"
        Image.new("RGB", (1200, 900), (60, 120, 90)).save(path)
        photos.append(InspectionPhoto(inspection_id="QA", section="PRODUCT_VIEW", file_path=path.name))
    document = Document()
    photo_grid(document, photos, Settings(data_dir=tmp_path), cols=5)
    row = [item for item in document.tables[0].rows if item._tr.xpath(".//wp:inline")][0]
    assert row.cells[0]._tc is row.cells[1]._tc
    assert row.cells[2]._tc is row.cells[3]._tc
    assert row.cells[4]._tc is not row.cells[0]._tc


def test_layout_patch_validation(tmp_path: Path) -> None:
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        inspection = create_basic_inspection(client)
        item_id, revision = inspection["id"], inspection["revision"]
        assert inspection["photo_layout"] == {}
        good = client.patch(f"/api/inspections/{item_id}", json={
            "revision": revision, "photo_layout": {"PRODUCT_VIEW": 3, "DIMENSION": 4}})
        assert good.status_code == 200, good.text
        assert good.json()["photo_layout"] == {"PRODUCT_VIEW": 3, "DIMENSION": 4}
        revision = good.json()["revision"]
        for bad in ({"PRODUCT_VIEW": 0}, {"PRODUCT_VIEW": 9}, {"NOPE": 2}):
            response = client.patch(f"/api/inspections/{item_id}", json={
                "revision": revision, "photo_layout": bad})
            assert response.status_code == 422, response.text
        null = client.patch(f"/api/inspections/{item_id}", json={"revision": revision, "photo_layout": None})
        assert null.status_code == 422, null.text


def test_layout_flows_into_docx_and_stays_deterministic(tmp_path: Path) -> None:
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        inspection = create_basic_inspection(client)
        item_id = inspection["id"]
        for index in range(6):
            response = client.post(f"/api/inspections/{item_id}/photos", data={"section": "PRODUCT_VIEW"},
                files=[("files", (f"view-{index}.jpg", image_bytes(), "image/jpeg"))])
            assert response.status_code == 201, response.text
        revision = client.get(f"/api/inspections/{item_id}").json()["revision"]
        assert client.patch(f"/api/inspections/{item_id}", json={
            "revision": revision, "photo_layout": {"PRODUCT_VIEW": 3}}).status_code == 200
        assert client.post(f"/api/inspections/{item_id}/validate").json()["errors"] == []
        first = client.post(f"/api/inspections/{item_id}/report/docx")
        assert first.status_code == 200, first.text
        word = Document(io.BytesIO(first.content))
        grid = next(table for table in word.tables
                    if len(table.columns) == 3 and table._tr.xpath(".//wp:inline"))
        photo_rows = [row for row in grid.rows if row._tr.xpath(".//wp:inline")]
        assert len(photo_rows) == 2
        second = client.post(f"/api/inspections/{item_id}/report/docx")
        assert second.content == first.content


def test_dimension_pair_follows_layout_columns(tmp_path: Path) -> None:
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        inspection = create_basic_inspection(client)
        item_id = inspection["id"]
        measurement = client.post(f"/api/inspections/{item_id}/measurements", json={
            "category": "DIMENSION", "label": "Length"}).json()
        response = client.post(f"/api/inspections/{item_id}/photos",
            data={"section": "DIMENSION", "measurement_id": measurement["id"]},
            files=[("files", (f"dim-{slot}.jpg", image_bytes(), "image/jpeg")) for slot in (1, 2)])
        assert response.status_code == 201, response.text
        revision = client.get(f"/api/inspections/{item_id}").json()["revision"]
        assert client.patch(f"/api/inspections/{item_id}", json={
            "revision": revision, "photo_layout": {"DIMENSION": 4}}).status_code == 200
        report = client.post(f"/api/inspections/{item_id}/report/docx")
        assert report.status_code == 200, report.text
        word = Document(io.BytesIO(report.content))
        table = next(table for table in word.tables if table.cell(0, 0).text.startswith("Length"))
        assert len(table.columns) == 4
        photo_row = [row for row in table.rows if row._tr.xpath(".//wp:inline")][0]
        assert photo_row.cells[0]._tc is photo_row.cells[1]._tc
        assert photo_row.cells[2]._tc is photo_row.cells[3]._tc
