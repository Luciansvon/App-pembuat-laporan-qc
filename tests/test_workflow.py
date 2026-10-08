"""A real inspection round trip through SQLite, photos and a Word document."""
from __future__ import annotations

import hashlib
import io
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from app.core.config import Settings
from app.main import create_app


def image_bytes(color: str = "#72584b") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (900, 600), color).save(buffer, format="JPEG")
    return buffer.getvalue()


def create_basic_inspection(client: TestClient, customer_code: str = "POLIFORM") -> dict:
    customer = client.post("/api/customers", json={"code": customer_code, "name": customer_code})
    assert customer.status_code == 201, customer.text
    product = client.post("/api/products", json={
        "customer_id": customer.json()["id"], "name": "QA FURNITURE MODEL",
        "standard_dimensions": {"L": 1200, "W": 600, "H": 450},
    })
    assert product.status_code == 201, product.text
    inspection = client.post("/api/inspections", json={
        "product_id": product.json()["id"], "po": "QA-PO-001,QA-PO-002",
        "inspection_type": "INLINE", "date": "2026-10-07", "location": "QSF",
        "qc_name": "Bima", "quantity": 20, "aql": "20", "box_dimensions": {},
    })
    assert inspection.status_code == 201, inspection.text
    return inspection.json()


def test_complete_inspection_generates_clean_word_and_persists(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path)
    app = create_app(settings)
    with TestClient(app) as client:
        inspection = create_basic_inspection(client)
        assert inspection["inspection_number"] == "QC-20261007-001"
        assert inspection["product_dimensions"] == {"L": 1200.0, "W": 600.0, "H": 450.0}
        item_id = inspection["id"]
        measurement = client.post(f"/api/inspections/{item_id}/measurements", json={
            "category": "DIMENSION", "label": "Height with glide", "value": 243,
            "unit": "mm", "standard_value": 223, "notes": "Measured with glide",
        })
        assert measurement.status_code == 201, measurement.text
        assert measurement.json()["result"] == "INFO"
        mc = client.post(f"/api/inspections/{item_id}/measurements", json={
            "category": "MC", "label": "Moisture Content", "unit": "%", "sort_order": 1})
        assert mc.status_code == 201
        for value in (9.4, 11.1):
            reading = client.post(f"/api/measurements/{mc.json()['id']}/readings", json={"value": value})
            assert reading.status_code == 201
        issue = client.post(f"/api/inspections/{item_id}/issues", json={
            "defect_type": "Gap", "quantity": 3,
            "cause": "During assembling press less strong", "repair": "Patch, small sanding",
            "cap": "Make sure before painting no gap on the joint"})
        assert issue.status_code == 201
        for photo_number, (section, linked_issue) in enumerate(
            (("PRODUCT_VIEW", None), *(("ISSUE", issue.json()["id"]) for _ in range(6)))):
            form = {"section": section}
            if linked_issue:
                form["issue_id"] = linked_issue
            response = client.post(f"/api/inspections/{item_id}/photos", data=form,
                files=[("files", (f"{section}.jpg", image_bytes(), "image/jpeg"))])
            assert response.status_code == 201, response.text
            assert client.get(f"/api/photos/{response.json()[0]['id']}/file").status_code == 200
            if section == "ISSUE" and photo_number in {1, 3, 5}:
                caption = client.patch(f"/api/photos/{response.json()[0]['id']}",
                    json={"caption": f"{500 + photo_number} mm (start from 100 mm)"})
                assert caption.status_code == 200
        review = client.post(f"/api/inspections/{item_id}/validate")
        assert review.status_code == 200
        assert review.json()["errors"] == []
        assert "Box dimension is empty" in review.json()["warnings"]
        report = client.post(f"/api/inspections/{item_id}/report/docx")
        assert report.status_code == 200, report.text
        assert report.content.startswith(b"PK")
        word = Document(io.BytesIO(report.content))
        header = " ".join(cell.text for row in word.sections[0].header.tables[0].rows for cell in row.cells)
        body = "\n".join(paragraph.text for paragraph in word.paragraphs)
        assert "QA-PO-001,QA-PO-002" in header
        assert "Bima" in header
        assert "{{" not in header
        assert "QC   INSPECTION REPORT" in header
        assert "Conclusion" in body
        assert "2 issue" not in body
        assert len(word.inline_shapes) == 7
        issue_table = next(table for table in word.tables if table.cell(0, 0).text == "ISSUE")
        assert len(issue_table.rows) == 5
        assert issue_table.cell(0, 1).text == "Gap ( 3 Pcs)"
        assert issue_table.cell(1, 0)._tc is issue_table.cell(1, 1)._tc
        assert len(issue_table.cell(1, 0)._tc.xpath('.//w:drawing')) == 6
        assert [issue_table.cell(index, 0).text for index in range(2, 5)] == ["CAUSE", "REPAIR", "CAP"]
        first_hash = hashlib.sha256(report.content).hexdigest()
        assert hashlib.sha256(client.post(f"/api/inspections/{item_id}/report/docx").content).hexdigest() == first_hash
        detail = client.get(f"/api/inspections/{item_id}").json()
        assert detail["measurements"][1]["value_min"] == 9.4
        assert detail["measurements"][1]["value_max"] == 11.1
        assert len(detail["photos"]) == 7
    with TestClient(create_app(settings)) as client:
        restored = client.get(f"/api/inspections/{item_id}")
        assert restored.status_code == 200
        assert restored.json()["po"] == "QA-PO-001,QA-PO-002"
        assert len(restored.json()["photos"]) == 7


def test_guards_invalid_photo_and_destructive_operations(tmp_path: Path) -> None:
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        inspection = create_basic_inspection(client)
        item_id = inspection["id"]
        invalid = client.post(f"/api/inspections/{item_id}/photos", data={"section": "PRODUCT_VIEW"},
            files=[("files", ("fake.jpg", b"not a photo", "image/jpeg"))])
        assert invalid.status_code == 422
        assert client.get(f"/api/inspections/{item_id}").json()["photos"] == []
        assert client.delete(f"/api/inspections/{item_id}").status_code == 428
        wrong_revision = client.patch(f"/api/inspections/{item_id}", json={"revision": 999, "po": "new"})
        assert wrong_revision.status_code == 409
        issue = client.post(f"/api/inspections/{item_id}/issues", json={"defect_type": "Gap"})
        assert issue.status_code == 201
        validation = client.post(f"/api/inspections/{item_id}/validate")
        assert any("has no photo" in error for error in validation.json()["errors"])
        assert client.post(f"/api/inspections/{item_id}/report/docx").status_code == 409


def test_single_issue_photo_and_caption_between_pairs(tmp_path: Path) -> None:
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        inspection = create_basic_inspection(client)
        item_id = inspection["id"]
        issue = client.post(f"/api/inspections/{item_id}/issues", json={
            "defect_type": "Gap", "cause": "Assembly pressure", "repair": "Patch",
            "cap": "Check the joint before painting"}).json()
        drawing = client.post(f"/api/inspections/{item_id}/photos",
            data={"section": "DRAWING"},
            files=[("files", ("drawing.jpg", image_bytes(), "image/jpeg"))])
        assert drawing.status_code == 201
        for order in range(2):
            response = client.post(f"/api/inspections/{item_id}/photos",
                data={"section": "DIMENSION", "sort_order": order},
                files=[("files", (f"dimension-{order}.jpg", image_bytes(), "image/jpeg"))])
            assert response.status_code == 201
            if order == 0:
                client.patch(f"/api/photos/{response.json()[0]['id']}", json={"caption": "503 mm (start from 100 mm)"})
        response = client.post(f"/api/inspections/{item_id}/photos",
            data={"section": "ISSUE", "issue_id": issue["id"]},
            files=[("files", ("gap.jpg", image_bytes(), "image/jpeg"))])
        assert response.status_code == 201
        report = client.post(f"/api/inspections/{item_id}/report/docx")
        assert report.status_code == 200, report.text
        word = Document(io.BytesIO(report.content))
        combined = next(table for table in word.tables if "Drawing Product" in table.cell(0, 0).text)
        assert "Product Dimension" in combined.cell(0, 0).text
        assert len(combined.cell(0, 0).tables) == 2
        dimension = combined.cell(0, 0).tables[1]
        assert "503 mm (start from 100 mm)" in dimension.cell(1, 0).text
        assert dimension.cell(1, 0)._tc is dimension.cell(1, 1)._tc
        issue_table = next(table for table in word.tables if table.cell(0, 0).text == "ISSUE")
        photo_row = issue_table.cell(1, 0).tables[0].rows[0]
        assert photo_row.cells[0]._tc is photo_row.cells[1]._tc
        assert len(photo_row.cells[0]._tc.xpath(".//w:drawing")) == 1
