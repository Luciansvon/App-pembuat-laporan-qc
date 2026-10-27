"""Template-backed deterministic DOCX report from confirmed database facts."""
from __future__ import annotations

import json
import os
import re
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile, ZipInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Emu, Pt, RGBColor
from fastapi import HTTPException
from PIL import Image
from sqlmodel import Session, select

from app.core.config import RESOURCE_ROOT, Settings
from app.models.tables import Inspection, InspectionPhoto, Issue, Measurement, TestReading
from app.report.photo_layouts import photo_layout_by_id
from app.services.storage import inspection_dir, photo_path

TEMPLATE_DIR = RESOURCE_ROOT / "templates"
SECTION_TITLES = {
    "PACKING_ATTRIBUTE": "ATRIBUTE( Box, packing ,Label, Position Logo)",
    "BOX_DIMENSION": "Box Dimension",
    "PACKING_METHOD": "Packing Method",
    "INNER_SKU_GLIDE": "Inner SKU & Fixed Glide",
    "PRODUCT_VIEW": "Product View", "PRODUCT_DETAIL": "Product Detail",
    "DRAIN_HOLE": "Drain Hole",
    "DRAWING": "Drawing Product", "DIMENSION": "Product Dimension",
    "LEVEL_CHECK": "Check with level",
    "MAL_CHECK": "Check With Mal",
    "RADIUS": "Check with Mal R 44",
    "RADIUS_R44": "Check with Mal R 44",
    "RADIUS_R10": "Check with Mal R 10",
    "MC": "MC Check", "INSPECTED_PRODUCT": "Inspected Product", "GLOSS": "Gloss Meter Check", "SWATCH": "Panel Swatch",
    "ISSUE": "ISSUE", "OTHER": "Other Inspection",
}


def section_title(config: dict, section: str) -> str:
    return config.get("section_titles", {}).get(section, SECTION_TITLES.get(section, section.replace("_", " ").title()))


class ReportAssetUnavailable(Exception):
    """A stored photo disappeared or became unreadable while building a report."""


def template_for(inspection: Inspection, settings: Settings) -> tuple[dict, bool]:
    catalog = json.loads((TEMPLATE_DIR / "catalog.json").read_text(encoding="utf-8"))
    report_type = getattr(inspection, "report_type", "STANDARD")
    if report_type == "COMBINED":
        key = "roda"
    elif report_type == "PACKING":
        packing_map = catalog.get("packing_templates", {})
        key = packing_map.get(inspection.customer_code.upper(), packing_map.get("DEFAULT", "rh_packing"))
    else:
        key = catalog["customer_templates"].get(inspection.customer_code.upper(), catalog["default_template"])
    config = catalog["templates"][key]
    path = TEMPLATE_DIR / config["output_template"]
    return config, config["status"] == "READY" and path.is_file()


def paragraph_format(paragraph, size: int = 9, bold: bool = False) -> None:
    paragraph.paragraph_format.space_after = Pt(4)
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(20, 27, 23)


def heading(document, label: str, page_break: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.page_break_before = page_break
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.add_run(label)
    paragraph_format(paragraph, 10, True)


def fill_header(document, inspection: Inspection, config: dict) -> None:
    def number(value: float | None) -> str:
        return "" if value is None else f"{value:g}"

    dimension_axes = config.get("dimension_axes", ["L", "W", "H"])
    box_axes = config.get("box_axes", ["L", "W", "H"])
    date_format = config.get("date_format", "%y%m%d")
    date_text = inspection.date.strftime(date_format)
    if config.get("date_label") == "DATE-LOC" and inspection.location:
        date_text += "-" + inspection.location
    inspect_text = (
        config.get("inline_label", "PACKING") if inspection.inspection_type == "PACKING"
        else config.get("inline_label", "INLINE") if inspection.inspection_type == "INLINE"
        else inspection.inspection_type.replace("_", " ")
    )
    include_product_dim = config.get("include_product_dim", True)
    values = {
        "DESC": inspection.product_name,
        "INSPECT": inspect_text,
        "CUSTOMER": inspection.customer_name, "PO": inspection.po,
        "DATE": date_text, "QC": inspection.qc_name,
        "QTY": str(inspection.quantity), "AQL": inspection.aql or "",
    }
    for index, axis in enumerate(dimension_axes, 1):
        values[f"DIM_{index}"] = number(inspection.product_dimensions.get(axis)) if include_product_dim else ""
    for index, axis in enumerate(box_axes, 1):
        values[f"BOX_{index}"] = number(inspection.box_dimensions.get(axis))

    def replace_placeholders(paragraph) -> bool:
        if "{{" not in paragraph.text:
            return False
        for run in paragraph.runs:
            if "{{" not in run.text:
                continue
            text = run.text
            for key, value in values.items():
                text = text.replace("{{" + key + "}}", value)
            run.text = text
        if "{{" in paragraph.text:
            # Placeholder terbelah antar-run: bangun ulang sekali sebagai fallback.
            replaced = paragraph.text
            for key, value in values.items():
                replaced = replaced.replace("{{" + key + "}}", value)
            paragraph.text = replaced
        return True

    for table in document.sections[0].header.tables:
        table.autofit = False
        target_twips = int(Cm(18).twips)
        grid_columns = list(table._tbl.tblGrid.gridCol_lst)
        grid_total = sum(column.w for column in grid_columns)
        if grid_total != target_twips and grid_total > 0:
            for column in grid_columns:
                column.w = int(column.w * target_twips / grid_total)
            for row in table.rows:
                for cell in row._tr.tc_lst:
                    cell.width = int(cell.width * Cm(18).emu / grid_total)
        table_width = table._tbl.tblPr.find(qn("w:tblW"))
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            table._tbl.tblPr.insert(0, table_width)
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(target_twips))
        for row_idx, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            trHeight = trPr.find(qn("w:trHeight"))
            if trHeight is None:
                trHeight = OxmlElement("w:trHeight")
                trPr.append(trHeight)
            if not trHeight.get(qn("w:val")):
                trHeight.set(qn("w:val"), "360" if row_idx == 0 else "320")
                trHeight.set(qn("w:hRule"), "exact")
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    if replace_placeholders(paragraph):
                        font_size = 8 if row_idx == len(table.rows) - 1 else 9
                        paragraph_format(paragraph, font_size)
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
    remaining = "\n".join(cell.text for table in document.sections[0].header.tables for row in table.rows for cell in row.cells)
    if "{{" in remaining:
        raise ValueError("DOCX header has unresolved fields")


def fitted_size(path: Path, max_width_cm: float = 7.5, max_height_cm: float = 6.0) -> tuple[float, float]:
    with Image.open(path) as image:
        width, height = image.size
    scale = min(max_width_cm / width, max_height_cm / height)
    return width * scale, height * scale


PHOTO_PADDING_CM = 0.12
PHOTO_PARENT_PADDING_CM = 0.20


def center_page_margins(document: Document) -> None:
    """Center the printable area without changing its total width."""
    for section in document.sections:
        side_margin_total = int(section.left_margin) + int(section.right_margin)
        centered_margin = Emu((side_margin_total + 1) // 2)
        section.left_margin = centered_margin
        section.right_margin = centered_margin


def photo_table_width(table, widths_cm: tuple[float, ...]) -> None:
    """Keep the grid, preferred width and cell padding in agreement for Word."""
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column, width in zip(table.columns, widths_cm, strict=True):
        column.width = Cm(width)
    properties = table._tbl.tblPr
    preferred = properties.find(qn("w:tblW"))
    if preferred is None:
        preferred = OxmlElement("w:tblW")
        properties.insert(0, preferred)
    preferred.set(qn("w:type"), "dxa")
    preferred.set(qn("w:w"), str(sum(Cm(width).twips for width in widths_cm)))
    margins = properties.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        properties.append(margins)
    for edge in ("top", "bottom", "left", "right"):
        margin = margins.find(qn(f"w:{edge}"))
        if margin is None:
            margin = OxmlElement(f"w:{edge}")
            margins.append(margin)
        margin.set(qn("w:type"), "dxa")
        margin.set(qn("w:w"), str(Cm(PHOTO_PADDING_CM).twips))


def photo_paragraph(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    # An inherited exact line height can make an inline image paint over a row border.
    paragraph.paragraph_format.line_spacing = 1.0


def cell_photo_width(cell) -> float:
    # Leave room for both explicit cell margins, including a merged final cell.
    return cell.width / Cm(1).emu - 2 * PHOTO_PADDING_CM


def insert_photo(paragraph, photo: InspectionPhoto, settings: Settings,
                 max_width_cm: float, max_height_cm: float) -> None:
    try:
        path = photo_path(settings, photo)
        width, height = fitted_size(path, max_width_cm=max_width_cm, max_height_cm=max_height_cm)
        paragraph.add_run().add_picture(str(path), width=Cm(width), height=Cm(height))
    except (HTTPException, OSError, ValueError, ArithmeticError, Image.DecompressionBombError) as exc:
        raise ReportAssetUnavailable(
            f"Photo {photo.id[:8]} ({photo.section}) became unavailable during report generation"
        ) from exc


def layout_cols(layout: dict | None, section: str) -> int:
    try:
        cols = int((layout or {}).get(section, 2))
    except (TypeError, ValueError, AttributeError):
        return 2
    return min(2, max(1, cols))


def row_spans(cols: int, count: int) -> list[int]:
    base, extra = divmod(cols, count)
    return [base + (1 if index < extra else 0) for index in range(count)]


def photo_grid(document, photos: list[InspectionPhoto], settings: Settings,
               *, width_cm: float = 18, border: bool = True,
               single_height_cm: float = 10, cols: int = 2,
               hide_captions: bool = False) -> None:
    if not photos:
        return
    cols = min(2, max(1, cols))
    table = document.add_table(rows=0, cols=cols)
    photo_table_width(table, tuple(width_cm / cols for _ in range(cols)))
    if border:
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "bottom", "left", "right"):
            edge_border = OxmlElement(f"w:{edge}")
            edge_border.set(qn("w:val"), "single")
            edge_border.set(qn("w:sz"), "4")
            borders.append(edge_border)
        for edge in ("insideH", "insideV"):
            edge_border = OxmlElement(f"w:{edge}")
            edge_border.set(qn("w:val"), "nil")
            borders.append(edge_border)
        table._tbl.tblPr.append(borders)
    for start in range(0, len(photos), cols):
        chunk = photos[start:start + cols]
        row = table.add_row()
        for cell in row.cells:
            cell.width = Cm(width_cm / cols)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        position = 0
        for photo, span in zip(chunk, row_spans(cols, len(chunk))):
            cell = row.cells[position]
            if span > 1:
                cell.merge(row.cells[position + span - 1])
            paragraph = cell.paragraphs[0]
            photo_paragraph(paragraph)
            insert_photo(paragraph, photo, settings,
                         cell_photo_width(cell),
                         single_height_cm if len(photos) == 1 else 5.8)
            position += span
        captions = [] if hide_captions else [photo.caption.strip() for photo in chunk if photo.caption and photo.caption.strip()]
        if captions:
            caption_cell = table.add_row().cells[0]
            if cols > 1:
                caption_cell = caption_cell.merge(table.rows[-1].cells[cols - 1])
            caption_cell.text = "  |  ".join(captions)
            caption = caption_cell.paragraphs[0]
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph_format(caption, 8, True)


def template_photo_grid(document, photos: list[InspectionPhoto | None], settings: Settings,
                        layout: dict, *, width_cm: float = 18) -> None:
    """Render the selected fixed-slot page while retaining one caption per photo."""
    columns = layout["columns"]
    widths = tuple(width_cm / columns for _ in range(columns))
    table = document.add_table(rows=0, cols=columns)
    photo_table_width(table, widths)
    margins = table._tbl.tblPr.find(qn("w:tblCellMar"))
    if margins is not None:
        for edge in ("top", "bottom"):
            margin = margins.find(qn(f"w:{edge}"))
            if margin is not None:
                margin.set(qn("w:w"), str(Cm(0.03).twips))
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "left", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        borders.append(border)
    for edge in ("insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for start in range(0, len(photos), columns):
        chunk = photos[start:start + columns]
        photo_row = table.add_row()
        caption_row = table.add_row()
        for cell in photo_row.cells:
            cell.width = Cm(width_cm / columns)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for cell in caption_row.cells:
            cell.width = Cm(width_cm / columns)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for index, photo in enumerate(chunk):
            paragraph = photo_row.cells[index].paragraphs[0]
            photo_paragraph(paragraph)
            paragraph.paragraph_format.keep_with_next = True
            if photo is not None:
                insert_photo(paragraph, photo, settings,
                             max_width_cm=cell_photo_width(photo_row.cells[index]),
                             max_height_cm=layout["max_photo_height_cm"])
            caption = caption_row.cells[index].paragraphs[0]
            caption.text = (photo.caption or "").strip() if photo is not None else ""
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph_format(caption, 8, True)
        for row in (photo_row, caption_row):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def dimension_photo_blocks(document, measurements: list[Measurement], photos: list[InspectionPhoto],
                           settings: Settings, *, width_cm: float = 18, border: bool = True,
                           cols: int = 2, label_position: str = "top") -> None:
    for measurement in measurements:
        pair = sorted((photo for photo in photos if photo.measurement_id == measurement.id),
                      key=lambda photo: (photo.dimension_slot or 0, photo.id))
        if len(pair) != 2 or {photo.dimension_slot for photo in pair} != {1, 2}:
            raise ReportAssetUnavailable(f"Dimension {measurement.label} requires exactly 2 linked photos")
        photo_grid(document, pair, settings, width_cm=width_cm, border=border, cols=cols)
        table = document.tables[-1]
        label_row = table.add_row()
        label_cell = label_row.cells[0]
        if cols > 1:
            label_cell = label_cell.merge(label_row.cells[cols - 1])
        label_cell.text = measurement.label
        if any(value is not None for value in (measurement.value, measurement.value_min, measurement.value_max)):
            label_cell.text += (" " if label_position == "bottom" else " — ") + value_text(measurement)
        paragraph_format(label_cell.paragraphs[0], 9, True)
        label_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if label_position == "top":
            table._tbl.insert(table._tbl.index(table.rows[0]._tr), label_row._tr)
        for index, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = index < len(table.rows) - 1


def combined_section_block(document, sections: list[str], photos_by_section: dict[str, list[InspectionPhoto]],
                           settings: Settings, dimensions: list[Measurement], layout: dict | None = None) -> None:
    layout = layout or {}
    outer = document.add_table(rows=1, cols=1)
    outer.style = "Table Grid"
    outer.autofit = False
    outer.columns[0].width = Cm(18)
    cell = outer.cell(0, 0)
    cell.width = Cm(18)
    for index, section in enumerate(sections):
        paragraph = cell.paragraphs[0] if index == 0 else cell.paragraphs[-1]
        if paragraph.text:
            paragraph = cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(8 if index else 5)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.add_run(SECTION_TITLES.get(section, section)).bold = True
        paragraph_format(paragraph, 9, True)
        if section == "DIMENSION":
            dimension_photo_blocks(cell, dimensions, photos_by_section[section], settings, width_cm=17.2, border=False,
                                   cols=layout_cols(layout, section))
        else:
            photo_grid(cell, photos_by_section[section], settings, width_cm=17.2,
                       border=False, single_height_cm=8.5, cols=layout_cols(layout, section))


def resolve_mc_range_text(mc_photos: list[InspectionPhoto], measurements: list[Measurement], session: Session) -> str | None:
    all_values: list[float] = []

    for photo in mc_photos:
        if photo.caption:
            caption_clean = photo.caption.strip()
            if "range mc" in caption_clean.lower():
                return caption_clean
            matches = re.findall(r"\b\d+(?:\.\d+)?\b", caption_clean)
            for m in matches:
                try:
                    val = float(m)
                    if 1.0 <= val <= 50.0:
                        all_values.append(val)
                except ValueError:
                    pass

    mc_measurements = [m for m in measurements if m.category == "MC"]
    for m in mc_measurements:
        if m.value is not None:
            all_values.append(float(m.value))
        if m.value_min is not None:
            all_values.append(float(m.value_min))
        if m.value_max is not None:
            all_values.append(float(m.value_max))
        readings = session.exec(select(TestReading).where(TestReading.measurement_id == m.id)).all()
        for r in readings:
            if r.value is not None:
                all_values.append(float(r.value))

    if all_values:
        min_val = min(all_values)
        max_val = max(all_values)
        return f"Range MC Check is {min_val:g}% - {max_val:g}%"
    return None


def value_text(measurement: Measurement) -> str:
    if measurement.value_min is not None or measurement.value_max is not None:
        low = measurement.value_min if measurement.value_min is not None else ""
        high = measurement.value_max if measurement.value_max is not None else ""
        return f"{low}–{high} {measurement.unit or ''}".strip()
    return f"{measurement.value if measurement.value is not None else '—'} {measurement.unit or ''}".strip()


def measurement_table(document, measurements: list[Measurement], session: Session) -> None:
    if not measurements:
        return
    heading(document, "Measurement and Test Results")
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ("CHECK", "ACTUAL", "STANDARD", "RESULT")):
        cell.text = label
    for measurement in measurements:
        cells = table.add_row().cells
        cells[0].text = f"{measurement.category} / {measurement.label}"
        cells[1].text = value_text(measurement)
        cells[2].text = str(measurement.standard_value) if measurement.standard_value is not None else "—"
        cells[3].text = measurement.result
        readings = session.exec(select(TestReading).where(TestReading.measurement_id == measurement.id)
            .order_by(TestReading.sort_order, TestReading.id)).all()
        if readings:
            paragraph = document.add_paragraph(f"{measurement.label} readings: " + ", ".join(str(reading.value) for reading in readings))
            paragraph_format(paragraph, 8)
        if measurement.notes:
            paragraph = document.add_paragraph(f"{measurement.label}: {measurement.notes}")
            paragraph_format(paragraph, 8)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph_format(paragraph, 8, row_index == 0)


def issue_block(document, issue: Issue, photos: list[InspectionPhoto], settings: Settings) -> None:
    # The reference uses one five-row table: issue, merged photo area, cause, repair, CAP.
    chunks = [photos[index:index + 6] for index in range(0, len(photos), 6)] or [[]]
    for index, chunk in enumerate(chunks):
        document.add_page_break()
        table = document.add_table(rows=5 if index == len(chunks) - 1 else 2, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        table.columns[0].width = Cm(2.44)
        table.columns[1].width = Cm(15.50)
        for row in table.rows:
            row.cells[0].width = Cm(2.44)
            row.cells[1].width = Cm(15.50)
        table.cell(0, 0).text = "ISSUE"
        title = issue.defect_type
        if issue.quantity is not None:
            title += f" ( {issue.quantity} Pcs)"
        if index:
            title += " — continued"
        table.cell(0, 1).text = title
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph_format(paragraph, 8)
        photo_cell = table.cell(1, 0).merge(table.cell(1, 1))
        if not chunk:
            photo_cell.text = "No issue photo recorded"
        else:
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "bottom", "left", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:type"), "dxa")
                margin.set(qn("w:w"), str(Cm(PHOTO_PARENT_PADDING_CM).twips))
                margins.append(margin)
            photo_cell._tc.get_or_add_tcPr().append(margins)
            grid = photo_cell.add_table(rows=0, cols=2)
            # Nested tables must fit inside the parent's content area, not its outer width.
            grid_width_cm = photo_cell.width / Cm(1).emu - 2 * PHOTO_PARENT_PADDING_CM
            photo_table_width(grid, (grid_width_cm / 2, grid_width_cm / 2))
            rows_count = (len(chunk) + 1) // 2
            # Share a bounded photo budget across 1, 2 or 3 rows, reserving room
            # for captions and CAUSE/REPAIR/CAP instead of growing with pixel size.
            max_h = min(9.5, 14.0 / rows_count)
            for pair_index in range(0, len(chunk), 2):
                pair = chunk[pair_index:pair_index + 2]
                row = grid.add_row()
                for cell in row.cells:
                    cell.width = Cm(grid_width_cm / 2)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if len(pair) == 1:
                    row.cells[0].merge(row.cells[1])
                for photo_index, photo in enumerate(pair):
                    cell = row.cells[photo_index]
                    paragraph = cell.paragraphs[0]
                    photo_paragraph(paragraph)
                    insert_photo(paragraph, photo, settings,
                                 max_width_cm=cell_photo_width(cell),
                                 max_height_cm=max_h)
                captions = [photo.caption.strip() for photo in pair if photo.caption and photo.caption.strip()]
                if captions:
                    caption_row = grid.add_row()
                    caption_cell = caption_row.cells[0].merge(caption_row.cells[1])
                    caption_cell.text = "  |  ".join(captions)
                    caption = caption_cell.paragraphs[0]
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph_format(caption, 8, True)
            for paragraph in photo_cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.line_spacing = Pt(1)
                paragraph.add_run().font.size = Pt(1)
            # Word needs the leading paragraph to separate a nested table from the
            # preceding ISSUE row. Removing it lets square images cover that border.
        if index != len(chunks) - 1:
            continue
        for row, (label, value) in zip(table.rows[2:], (
            ("CAUSE", issue.cause or "Not recorded"),
            ("REPAIR", issue.repair or "Not recorded"),
            ("CAP", issue.cap or "Not recorded"),
        )):
            row.cells[0].text = label
            row.cells[1].text = value
            for cell_index, cell in enumerate(row.cells):
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = label != "CAP"
                    paragraph_format(paragraph, 8, cell_index == 0)


def canonicalize_zip(path: Path) -> None:
    temporary = path.with_suffix(".canonical.tmp")
    with ZipFile(path) as source, ZipFile(temporary, "w") as output:
        for member in source.infolist():
            info = ZipInfo(member.filename, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = member.compress_type
            info.external_attr = member.external_attr
            output.writestr(info, source.read(member.filename))
    os.replace(temporary, path)


def report_filename(inspection: Inspection) -> str:
    def slug(value: str) -> str:
        return re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_")[:70] or "UNKNOWN"
    return f"{slug(inspection.customer_name)}-{slug(inspection.product_name)}-{inspection.date:%y%m%d}-{inspection.inspection_type}-{inspection.inspection_number}-r{inspection.revision}.docx"


def generate_report(session: Session, inspection: Inspection, settings: Settings) -> Path:
    config, ready = template_for(inspection, settings)
    if not ready:
        raise ValueError("DOCX template is not ready")
    document = Document(TEMPLATE_DIR / config["output_template"])
    center_page_margins(document)
    if document.paragraphs and not document.paragraphs[0].text.strip():
        p_el = document.paragraphs[0]._element
        p_el.getparent().remove(p_el)
    fill_header(document, inspection, config)
    photos = session.exec(select(InspectionPhoto).where(InspectionPhoto.inspection_id == inspection.id)
        .order_by(InspectionPhoto.sort_order, InspectionPhoto.id)).all()
    measurements = session.exec(select(Measurement).where(Measurement.inspection_id == inspection.id)
        .order_by(Measurement.sort_order, Measurement.id)).all()
    dimensions = [measurement for measurement in measurements if measurement.category in ("DIMENSION", "BOX_DIMENSION")]
    issues = session.exec(select(Issue).where(Issue.inspection_id == inspection.id)
        .order_by(Issue.sort_order, Issue.id)).all()
    assigned_photo_ids = {photo_id for page in inspection.photo_pages or []
                          for photo_id in page.get("photo_ids", []) if photo_id}
    sections_photos = {section: [photo for photo in photos if photo.section == section
                                 and photo.issue_id is None and photo.id not in assigned_photo_ids]
                       for section in config["sections"]}
    combined = {block["sections"][0]: block for block in config.get("combined_blocks", [])}
    break_before = set(config.get("break_before", []))
    layout = inspection.photo_layout or {}
    consumed: set[str] = set()
    started = False
    for section in config["sections"]:
        if section in {"ISSUE", "CONCLUSION"} or section in consumed:
            continue
        section_photos = sections_photos[section]
        if not section_photos:
            continue
        block = combined.get(section)
        if block and all(0 < len(sections_photos.get(name, [])) <= limit for name, limit
                         in zip(block["sections"], block["max_counts"], strict=True)):
            if started and any(name in break_before for name in block["sections"]):
                document.add_page_break()
            combined_section_block(document, block["sections"], sections_photos, settings, dimensions, layout)
            consumed.update(block["sections"])
            started = True
            continue
        if section in ("DIMENSION", "BOX_DIMENSION"):
            target_measurements = [m for m in measurements if m.category == section]
            if target_measurements:
                heading(document, section_title(config, section), page_break=started and section in break_before)
                dimension_photo_blocks(document, target_measurements, section_photos, settings,
                                       cols=layout_cols(layout, section),
                                       label_position=config.get("dimension_label_position", "top"))
                started = True
            continue
        if section == "MC":
            heading(document, section_title(config, section),
                    page_break=started and section in break_before)
            range_text = resolve_mc_range_text(section_photos, measurements, session)
            only_numbers = all(
                not p.caption or bool(re.match(r"^[\d\.\s%–\-]+$", p.caption.strip()))
                for p in section_photos
            )
            photo_grid(document, section_photos, settings,
                       cols=layout_cols(layout, section),
                       hide_captions=bool(range_text and only_numbers))
            if range_text:
                p_range = document.add_paragraph(range_text)
                p_range.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_range.paragraph_format.space_before = Pt(4)
                p_range.paragraph_format.space_after = Pt(8)
                paragraph_format(p_range, 9, bold=True)
            started = True
            continue

        if section == "INSPECTED_PRODUCT":
            qty = inspection.inspected_quantity or inspection.quantity or 1
            for p in section_photos:
                if not p.caption or not p.caption.strip():
                    p.caption = f"inspect {qty} pcs"
                elif not p.caption.lower().startswith("inspect") and "pcs" not in p.caption.lower():
                    p.caption = f"inspect {p.caption.strip()} pcs"
            photo_grid(document, section_photos, settings,
                       cols=1, single_height_cm=9.5)
            started = True
            continue

        heading(document, section_title(config, section),
                page_break=started and section in break_before)
        photo_grid(document, section_photos, settings,
                   cols=layout_cols(layout, section))
        started = True
    photos_by_id = {photo.id: photo for photo in photos}
    for index, page in enumerate(inspection.photo_pages or [], start=1):
        layout_config = photo_layout_by_id(page["template_id"])
        if layout_config is None:
            raise ReportAssetUnavailable("Selected photo page layout is unavailable")
        page_photos = [photos_by_id.get(photo_id) for photo_id in page["photo_ids"]]
        if started:
            document.add_page_break()
        title = section_title(config, page["section"])
        heading(document, f"{title} · {layout_config['code']} · Page {index}")
        template_photo_grid(document, page_photos, settings, layout_config)
        started = True
    if config.get("include_measurement_table", False):
        measurement_table(document, measurements, session)
    for issue in issues:
        issue_block(document, issue, [photo for photo in photos if photo.issue_id == issue.id], settings)
    if config.get("include_conclusion", True):
        heading(document, "Conclusion", page_break=False)
        conclusion = inspection.conclusion or f"Inspection recorded. {len(issues)} issue(s) documented."
        paragraph = document.add_paragraph(conclusion)
        paragraph_format(paragraph, 9)
        if inspection.communication_status:
            paragraph = document.add_paragraph(inspection.communication_status)
            paragraph_format(paragraph, 9)
    fixed_time = datetime(2000, 1, 1)
    document.core_properties.created = fixed_time
    document.core_properties.modified = fixed_time
    destination = inspection_dir(settings, inspection) / "output" / report_filename(inspection)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name("." + destination.name + ".tmp")
    try:
        document.save(temporary)
        canonicalize_zip(temporary)
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)
    return destination
