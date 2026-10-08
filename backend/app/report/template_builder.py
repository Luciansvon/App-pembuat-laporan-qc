"""Make clean, customer-specific Word templates from inspected originals.

The source body and all body image relationships are removed before saving.
Run only on privately retained references with matching hashes.
"""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
from zipfile import ZipFile, ZipInfo

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[3]
TEMPLATES = ROOT / "templates"


def remove_package_thumbnail(target: Path) -> None:
    """Word may copy a large visual preview of the historical report."""
    temporary = target.with_suffix(".clean.tmp")
    with ZipFile(target) as source, ZipFile(temporary, "w") as output:
        for member in source.infolist():
            if member.filename.startswith("docProps/thumbnail."):
                continue
            payload = source.read(member.filename)
            if member.filename == "_rels/.rels":
                root = etree.fromstring(payload)
                for relationship in list(root):
                    if relationship.get("Type", "").endswith("/thumbnail"):
                        root.remove(relationship)
                payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif member.filename == "[Content_Types].xml":
                root = etree.fromstring(payload)
                for part in list(root):
                    if "thumbnail" in part.get("PartName", ""):
                        root.remove(part)
                payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            info = ZipInfo(member.filename, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = member.compress_type
            info.external_attr = member.external_attr
            output.writestr(info, payload)
    os.replace(temporary, target)


def set_slot(cell, token: str) -> None:
    cell.text = "{{" + token + "}}"
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8)


def build(name: str, record: dict) -> Path:
    source = TEMPLATES / record["reference"]
    if hashlib.sha256(source.read_bytes()).hexdigest() != record["reference_sha256"]:
        raise ValueError(f"Reference hash mismatch: {name}")
    document = Document(source)
    body = document.element.body
    for element in list(body):
        if element.tag != qn("w:sectPr"):
            body.remove(element)
    for relationship_id, relationship in list(document.part.rels.items()):
        if relationship.reltype.endswith("/image"):
            document.part.drop_rel(relationship_id)
    document.add_paragraph()
    header = document.sections[0].header.tables[0]
    header.autofit = False
    slots = {
        (1, 1): "DESC", (1, 13): "INSPECT", (2, 1): "CUSTOMER", (2, 9): "PO",
        (3, 1): "DATE", (3, 9): "QC", (4, 1): "QTY", (4, 9): "AQL",
        (5, 2): "DIM_1", (5, 4): "DIM_2", (5, 6): "DIM_3",
        (5, 9): "BOX_1", (5, 11): "BOX_2", (5, 14): "BOX_3",
    }
    original_fields = {header.rows[row].cells[column].text.strip() for row, column in slots}
    for (row, column), token in slots.items():
        set_slot(header.rows[row].cells[column], token)
    for row, column in ((2, 0), (5, 7)):
        for paragraph in header.rows[row].cells[column].paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(8)
    # The retained last row expands to the full header height after source photos
    # are removed. Its values are bounded short dimensions, so keep that row compact.
    header.rows[5].height = Pt(16)
    header.rows[5].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    document.core_properties.author = "Inspectra"
    document.core_properties.last_modified_by = "Inspectra"
    document.core_properties.title = "QC Inspection Report Template"
    document.core_properties.comments = ""
    document.core_properties.keywords = ""
    target = TEMPLATES / record["output_template"]
    document.save(target)
    remove_package_thumbnail(target)
    with ZipFile(target) as archive:
        if any(part.startswith("word/media/") or part.startswith("docProps/thumbnail.") for part in archive.namelist()):
            raise AssertionError("Template still contains source visuals")
        content = b"".join(archive.read(part) for part in archive.namelist() if part.endswith(".xml"))
        for original in (value.encode("utf-8") for value in original_fields if len(value) >= 5):
            if original in content:
                raise AssertionError(f"Historical field remains in {target.name}")
    return target


if __name__ == "__main__":
    catalog = json.loads((TEMPLATES / "catalog.json").read_text(encoding="utf-8"))
    for key in ("poliform", "rh", "default"):
        result = build(key, catalog["templates"][key])
        print(f"Clean template: {result}")
